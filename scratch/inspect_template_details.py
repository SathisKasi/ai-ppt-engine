import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation

prs = Presentation('AITransformationWeeklyUpdate4SEP2026.pptx')
print(f"Presentation dimensions: {prs.slide_width} x {prs.slide_height}")

for i, s in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ({s.slide_layout.name}) ====================")
    if s.background and s.background.fill:
        print(f"  Slide BG Fill Type: {s.background.fill.type}")
    for sh in s.shapes:
        txt = sh.text.replace('\n', ' ')[:60] if hasattr(sh, 'text') else ''
        print(f"  Shape {sh.shape_id:3}: {sh.name:25} | type={sh.shape_type} | pos=({sh.left}, {sh.top}, {sh.width}, {sh.height}) | '{txt}'")
        if sh.name == 'Group 8':
            print("    Group 8 sub-shapes:")
            for sub in sh.shapes:
                sub_txt = sub.text.replace('\n', ' ')[:40] if hasattr(sub, 'text') else ''
                print(f"      SubShape {sub.shape_id:3}: {sub.name:20} | '{sub_txt}'")
