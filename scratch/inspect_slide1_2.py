from pptx import Presentation

prs = Presentation('AITransformationWeeklyUpdate4SEP2026.pptx')
for idx in [0, 1]:
    s = prs.slides[idx]
    print(f"\n=== SLIDE {idx+1} ({s.slide_layout.name}) ===")
    for sh in s.shapes:
        txt = sh.text.replace('\n', ' ') if hasattr(sh, 'text') else ''
        print(f"  Shape {sh.shape_id}: {sh.name} | type={sh.shape_type} | pos=({sh.left}, {sh.top}, {sh.width}, {sh.height}) | '{txt}'")
        if sh.name == 'Group 8':
            for sub in sh.shapes:
                st = sub.text.replace('\n', ' ') if hasattr(sub, 'text') else ''
                print(f"    Sub: {sub.name} pos=({sub.left}, {sub.top}, {sub.width}, {sub.height}) | '{st}'")
