from pptx import Presentation

prs = Presentation('AITransformationWeeklyUpdate4SEP2026.pptx')
s3 = prs.slides[2]
print("--- SLIDE 3 CARD SHAPES ---")
for s in s3.shapes:
    if "Shape" in s.name and s.width > 2000000:
        fill_type = s.fill.type if s.fill else None
        rgb = None
        try:
            rgb = s.fill.fore_color.rgb
        except Exception:
            pass
        print(f"  {s.name} ({s.shape_type}) pos=({s.left},{s.top},{s.width},{s.height}) fill={fill_type} rgb={rgb}")

s4 = prs.slides[3]
print("\n--- SLIDE 4 CARD SHAPES ---")
for s in s4.shapes:
    if "Shape" in s.name and s.width > 2000000:
        fill_type = s.fill.type if s.fill else None
        rgb = None
        try:
            rgb = s.fill.fore_color.rgb
        except Exception:
            pass
        print(f"  {s.name} ({s.shape_type}) pos=({s.left},{s.top},{s.width},{s.height}) fill={fill_type} rgb={rgb}")
