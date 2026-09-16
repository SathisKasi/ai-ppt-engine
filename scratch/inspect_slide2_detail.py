from pptx import Presentation

prs = Presentation('AITransformationWeeklyUpdate4SEP2026.pptx')
s2 = prs.slides[1]
print(f"Slide 2 shapes ({len(s2.shapes)}):")
for sh in s2.shapes:
    print(f"  {sh.name} ({sh.shape_type}) pos=({sh.left},{sh.top},{sh.width},{sh.height})")
    if sh.name == 'Group 8':
        for sub in sh.shapes:
            print(f"    Sub: {sub.name} ({sub.shape_type}) pos=({sub.left},{sub.top},{sub.width},{sub.height})")
            if hasattr(sub, 'text_frame'):
                for p_idx, p in enumerate(sub.text_frame.paragraphs):
                    print(f"      P{p_idx}: text='{p.text}' font_size={p.font.size}")
                    for r_idx, r in enumerate(p.runs):
                        print(f"        R{r_idx}: text='{r.text}' font_name={r.font.name} font_size={r.font.size} bold={r.font.bold}")
