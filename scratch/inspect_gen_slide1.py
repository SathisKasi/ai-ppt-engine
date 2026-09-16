from pptx import Presentation

prs = Presentation('output/test_template1_user_plan_fixed.pptx')
s1 = prs.slides[0]
for sh in s1.shapes:
    print(f"Shape {sh.shape_id}: {sh.name} pos=({sh.left},{sh.top},{sh.width},{sh.height})")
    if hasattr(sh, 'text_frame'):
        for p in sh.text_frame.paragraphs:
            print(f"  P: font_size={p.font.size} text='{p.text}'")
            for r in p.runs:
                print(f"    R: font_size={r.font.size} text='{r.text}'")
