from pptx import Presentation

prs = Presentation('AITransformationWeeklyUpdate4SEP2026.pptx')
s1 = prs.slides[0]
for sh in s1.shapes:
    print(f"\nShape: {sh.name} ({sh.shape_type}) pos=({sh.left},{sh.top},{sh.width},{sh.height})")
    if hasattr(sh, 'text_frame'):
        tf = sh.text_frame
        for p_idx, p in enumerate(tf.paragraphs):
            print(f"  P{p_idx}: text='{p.text}'")
            for r_idx, r in enumerate(p.runs):
                print(f"    R{r_idx}: font_name={r.font.name} font_size={r.font.size} bold={r.font.bold} text='{r.text}'")
