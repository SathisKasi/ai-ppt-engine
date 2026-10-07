import json
from pathlib import Path
from pptx import Presentation
from core.builders.hld_qbr_generic_builder import HLDQBRGenericBuilder
from llm.hld_qbr_generic_schemas import GenericHLDQBRPlan

plan_path = Path("logs/llm/20261007_162734_90e2f3/presentation_plan.json")
plan_dict = json.loads(plan_path.read_text(encoding="utf-8"))
plan = GenericHLDQBRPlan(**plan_dict)

print(f"Loaded plan with {len(plan.slides)} content slides.")

builder = HLDQBRGenericBuilder()
pptx_bytes = builder.build(plan)

out_test = Path("output/test_fixed_deck_verification.pptx")
out_test.write_bytes(pptx_bytes)
print(f"Saved test output to {out_test} ({len(pptx_bytes)} bytes)")

prs = Presentation(str(out_test))
print(f"Generated deck has {len(prs.slides)} slides:")

for idx, slide in enumerate(prs.slides, 1):
    titles = []
    badges = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip().replace("\n", " ")
            if "slide number" in shape.name.lower():
                badges.append(txt)
            elif txt:
                titles.append(txt[:50])
    main_text = titles[0] if titles else "EMPTY"
    badge_str = badges[0] if badges else "None"
    print(f"  Slide {idx:2d}: badge='{badge_str}' | first_text='{main_text}'")
