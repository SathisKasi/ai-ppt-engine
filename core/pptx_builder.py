"""
core/pptx_builder.py — PowerPoint generation engine.

ROOT CAUSE FIX: Copying XML from template slides carries layout inheritance,
which causes python-pptx to revert to placeholder text on save+reload.

SOLUTION: Each slide is drawn FROM SCRATCH using python-pptx shape primitives
(add_shape, add_textbox) with the same design as create_template.py,
but populated with LLM-generated content.

This guarantees:
  - Content always persists after save/reload
  - All shapes remain fully editable in PowerPoint
  - Design is consistent with the template aesthetic
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from llm.schemas import PresentationPlan, SlideDefinition
from core.layout_manager import LayoutManager
from utils.logging_utils import get_logger

logger = get_logger(__name__)

# ---------------------------------------------------------------------------
# Slide dimensions
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ---------------------------------------------------------------------------
# Color palette (matches create_template.py)
# ---------------------------------------------------------------------------
BG_DARK    = RGBColor(0x1A, 0x1A, 0x2E)   # Deep navy
BG_CARD    = RGBColor(0x16, 0x21, 0x3E)   # Card background
ACCENT_BLUE   = RGBColor(0x0F, 0x34, 0x60)
ACCENT_PURPLE = RGBColor(0x53, 0x34, 0x83)
ACCENT_RED    = RGBColor(0xE9, 0x45, 0x60)
ACCENT_TEAL   = RGBColor(0x00, 0x88, 0xCC)
ACCENT_GREEN  = RGBColor(0x00, 0x99, 0x66)
TEXT_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_LIGHT = RGBColor(0xE0, 0xE0, 0xE0)
TEXT_MUTED = RGBColor(0xA0, 0xA0, 0xB0)

FONT = "Calibri"

# ---------------------------------------------------------------------------
# Low-level drawing helpers
# ---------------------------------------------------------------------------

def _bg(slide, color: RGBColor = BG_DARK) -> None:
    """Fill slide background."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left, top, width, height,
          fill: RGBColor = ACCENT_BLUE, line: bool = False) -> Any:
    """Add a filled rectangle."""
    shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if not line:
        shape.line.fill.background()
    return shape


def _oval(slide, left, top, width, height, fill: RGBColor = ACCENT_RED) -> Any:
    """Add a filled oval."""
    shape = slide.shapes.add_shape(9, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _txt(slide, text: str, left: float, top: float, width: float, height: float,
         size: int = 18, bold: bool = False, color: RGBColor = TEXT_LIGHT,
         align=PP_ALIGN.LEFT, wrap: bool = True) -> Any:
    """Add a styled textbox. Returns the shape."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return txBox


def _bullets(slide, items: List[str], left: float, top: float, width: float, height: float,
             size: int = 17, color: RGBColor = TEXT_LIGHT, char: str = "• ") -> Any:
    """Add a multi-bullet textbox. Returns the shape."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        cleaned = re.sub(r"^[\s•\-\*\u2022\u2023\u25aa]+", "", str(item)).strip()
        if not cleaned:
            continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = f"{char}{cleaned}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = FONT
    return txBox


def _header_bar(slide, title: str, title_size: int = 28) -> None:
    """Draw the standard dark header bar with title."""
    _rect(slide, 0, 0, 13.333, 1.4, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.08,  1.4, ACCENT_RED)
    _txt(slide, title, 0.3, 0.1, 12.5, 1.1,
         size=title_size, bold=True, color=TEXT_WHITE)


def _ensure_list(val: Any) -> List[str]:
    if isinstance(val, list):
        return [str(v) for v in val if v]
    if isinstance(val, str) and val.strip():
        return [val.strip()]
    return []


# ---------------------------------------------------------------------------
# Layout drawers — each builds a complete slide from scratch
# ---------------------------------------------------------------------------

def _draw_title(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """TITLE — cover slide."""
    _bg(slide, BG_DARK)
    _rect(slide, 0, 6.2, 13.333, 0.1, ACCENT_RED)
    _rect(slide, 0, 0,   13.333, 0.08, ACCENT_PURPLE)
    _rect(slide, 0, 0,   0.05,  7.5,  ACCENT_PURPLE)

    _txt(slide, slide_def.title,
         1.0, 2.0, 11.0, 1.5,
         size=44, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

    subtitle = content.get("subtitle", "")
    presenter = content.get("presenter", "")
    date_str  = content.get("date", "")
    parts = [p for p in [subtitle, presenter, date_str] if p]
    sub_text = " • ".join(parts) if parts else slide_def.purpose

    _txt(slide, sub_text,
         1.0, 3.8, 11.0, 0.8,
         size=22, bold=False, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)

    _txt(slide, "CONFIDENTIAL", 0.3, 6.8, 3.0, 0.4,
         size=10, color=TEXT_MUTED)


def _draw_executive_summary(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """EXECUTIVE_SUMMARY — Slide 2: key highlights for decision-makers."""
    _bg(slide, BG_DARK)

    # Header band
    _rect(slide, 0, 0, 13.333, 1.5, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.1, 1.5, ACCENT_RED)
    _txt(slide, slide_def.title or "Executive Summary",
         0.3, 0.1, 10.0, 0.85,
         size=28, bold=True, color=TEXT_WHITE)
    _txt(slide, "KEY HIGHLIGHTS",
         0.3, 0.92, 4.0, 0.4,
         size=11, bold=False, color=ACCENT_RED)

    highlights = _ensure_list(content.get("highlights", []))
    purpose   = content.get("purpose_statement", "")
    key_metric = content.get("key_metric", "")

    # Highlight cards (up to 5)
    accent_colors = [ACCENT_RED, ACCENT_PURPLE, ACCENT_TEAL, ACCENT_GREEN, ACCENT_BLUE]
    card_h = 0.9
    start_y = 1.7
    for i, hl in enumerate(highlights[:5]):
        color = accent_colors[i % len(accent_colors)]
        y = start_y + i * (card_h + 0.12)
        # Left accent strip
        _rect(slide, 0.3, y, 0.06, card_h, color)
        # Card background
        _rect(slide, 0.4, y, 11.8, card_h, BG_CARD)
        # Number badge
        _txt(slide, str(i + 1), 0.55, y + 0.18, 0.4, 0.5,
             size=16, bold=True, color=color)
        # Highlight text
        _txt(slide, hl, 1.1, y + 0.14, 11.0, card_h - 0.1,
             size=17, bold=False, color=TEXT_LIGHT)

    # Purpose statement at bottom
    if purpose:
        bottom_y = start_y + min(len(highlights), 5) * (card_h + 0.12) + 0.12
        if bottom_y < 6.8:
            _rect(slide, 0.3, bottom_y, 12.7, 0.55, ACCENT_PURPLE)
            _txt(slide, purpose, 0.55, bottom_y + 0.08, 12.2, 0.4,
                 size=14, bold=False, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

    # Key metric callout (top-right corner)
    if key_metric:
        _rect(slide, 10.5, 0.05, 2.75, 1.35, BG_CARD)
        _txt(slide, key_metric, 10.55, 0.1, 2.65, 1.25,
             size=18, bold=True, color=ACCENT_RED, align=PP_ALIGN.CENTER)


def _draw_agenda(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """AGENDA — Slide 3: numbered list of major presentation topics."""
    _bg(slide, BG_DARK)

    # Header band
    _rect(slide, 0, 0, 13.333, 1.4, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.1, 1.4, ACCENT_RED)
    _txt(slide, slide_def.title or "Agenda",
         0.3, 0.1, 10.0, 0.8,
         size=28, bold=True, color=TEXT_WHITE)
    _txt(slide, "PRESENTATION OVERVIEW",
         0.3, 0.85, 5.0, 0.35,
         size=11, bold=False, color=ACCENT_RED)

    agenda_items = _ensure_list(content.get("agenda_items", []))
    descriptions = _ensure_list(content.get("descriptions", []))
    time_alloc   = _ensure_list(content.get("time_allocation", []))

    # Layout: two columns if >4 items, single column otherwise
    n = len(agenda_items)
    accent_colors = [ACCENT_RED, ACCENT_PURPLE, ACCENT_TEAL, ACCENT_GREEN,
                     ACCENT_BLUE, ACCENT_RED, ACCENT_PURPLE, ACCENT_TEAL]

    if n <= 4:
        # Single wide column
        item_h = 5.5 / max(n, 1)
        for i, item in enumerate(agenda_items[:8]):
            color = accent_colors[i % len(accent_colors)]
            y = 1.6 + i * item_h
            desc = descriptions[i] if i < len(descriptions) else ""
            time = time_alloc[i] if i < len(time_alloc) else ""

            _rect(slide, 0.4, y + 0.05, 12.5, item_h - 0.1, BG_CARD)
            # Number circle
            _oval(slide, 0.55, y + 0.12, item_h * 0.55, item_h * 0.55, color)
            _txt(slide, str(i + 1), 0.55, y + 0.15, item_h * 0.55, item_h * 0.5,
                 size=18, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
            # Item title
            title_x = 0.55 + item_h * 0.6
            _txt(slide, item, title_x, y + 0.12, 10.5, 0.5,
                 size=17, bold=True, color=TEXT_WHITE)
            if desc:
                _txt(slide, desc, title_x, y + 0.55, 10.0, 0.35,
                     size=13, color=TEXT_MUTED)
            if time:
                _txt(slide, time, 11.8, y + 0.15, 1.0, 0.4,
                     size=12, bold=True, color=color, align=PP_ALIGN.RIGHT)
    else:
        # Two-column layout for 5-8 items
        left_items  = agenda_items[:4]
        right_items = agenda_items[4:8]
        for col, items_col in enumerate([left_items, right_items]):
            x_base = 0.4 if col == 0 else 6.9
            col_w  = 6.0
            item_h = 5.6 / max(len(items_col), 1)
            for i, item in enumerate(items_col):
                global_i = i if col == 0 else i + len(left_items)
                color = accent_colors[global_i % len(accent_colors)]
                y = 1.6 + i * item_h
                desc = descriptions[global_i] if global_i < len(descriptions) else ""

                _rect(slide, x_base, y + 0.05, col_w, item_h - 0.1, BG_CARD)
                _oval(slide, x_base + 0.15, y + 0.12, 0.5, 0.5, color)
                _txt(slide, str(global_i + 1), x_base + 0.15, y + 0.14, 0.5, 0.45,
                     size=14, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
                _txt(slide, item, x_base + 0.8, y + 0.12, col_w - 0.9, 0.5,
                     size=15, bold=True, color=TEXT_WHITE)
                if desc:
                    _txt(slide, desc, x_base + 0.8, y + 0.55, col_w - 0.9, 0.35,
                         size=12, color=TEXT_MUTED)


def _draw_title_content(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """TITLE_AND_CONTENT — title + bullet points."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)
    _rect(slide, 0.3, 1.6, 12.7, 5.5, BG_CARD)

    bullets   = _ensure_list(content.get("bullets", []))
    body_text = content.get("body_text", "")

    if bullets:
        _bullets(slide, bullets[:8], 0.6, 1.8, 12.2, 5.0, size=18, color=TEXT_LIGHT)
    elif body_text:
        _txt(slide, body_text, 0.6, 1.8, 12.2, 5.0, size=17, color=TEXT_LIGHT)
    else:
        _txt(slide, slide_def.purpose or "Content", 0.6, 1.8, 12.2, 5.0,
             size=17, color=TEXT_MUTED)


def _draw_two_column(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """TWO_COLUMN — side-by-side layout."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    left_heading  = content.get("left_heading",  "Left Column")
    right_heading = content.get("right_heading", "Right Column")
    left_bullets  = _ensure_list(content.get("left_bullets",  []))
    right_bullets = _ensure_list(content.get("right_bullets", []))

    # Left column
    _rect(slide, 0.3, 1.6, 6.1, 5.5, BG_CARD)
    _txt(slide,  left_heading, 0.4, 1.7, 5.8, 0.6,
         size=16, bold=True, color=ACCENT_RED)
    _bullets(slide, left_bullets[:6], 0.4, 2.4, 5.8, 4.4, size=16, color=TEXT_LIGHT)

    # Divider
    _rect(slide, 6.6, 1.6, 0.05, 5.5, ACCENT_PURPLE)

    # Right column
    _rect(slide, 6.9, 1.6, 6.1, 5.5, BG_CARD)
    _txt(slide,  right_heading, 7.0, 1.7, 5.8, 0.6,
         size=16, bold=True, color=ACCENT_PURPLE)
    _bullets(slide, right_bullets[:6], 7.0, 2.4, 5.8, 4.4, size=16, color=TEXT_LIGHT)


def _draw_section_header(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """SECTION_HEADER — chapter divider."""
    _bg(slide, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.08, 7.5, ACCENT_RED)
    _rect(slide, 0, 3.5, 13.333, 0.06, ACCENT_RED)

    section_title = content.get("section_title", slide_def.title)
    description   = content.get("description", slide_def.purpose or "")

    _txt(slide, section_title,
         1.0, 1.8, 11.0, 1.4,
         size=40, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

    if description:
        _txt(slide, description,
             1.5, 3.7, 10.0, 0.8,
             size=18, bold=False, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)


def _draw_image_text(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """IMAGE_TEXT — image placeholder left, text right."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    # Image placeholder
    _rect(slide, 0.3, 1.6, 6.0, 5.5, BG_CARD)
    _txt(slide, "[ Image / Diagram ]",
         0.5, 3.5, 5.6, 1.0, size=16, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    # Text area
    _rect(slide, 6.6, 1.6, 6.4, 5.5, BG_CARD)

    description = content.get("description", "")
    bullets     = _ensure_list(content.get("bullets", []))

    if bullets:
        if description:
            _txt(slide, description, 6.8, 1.8, 6.0, 0.8, size=15, color=TEXT_LIGHT)
            _bullets(slide, bullets[:5], 6.8, 2.7, 6.0, 4.0, size=15, color=TEXT_LIGHT)
        else:
            _bullets(slide, bullets[:5], 6.8, 1.8, 6.0, 5.0, size=16, color=TEXT_LIGHT)
    elif description:
        _txt(slide, description, 6.8, 1.8, 6.0, 5.0, size=16, color=TEXT_LIGHT)


def _draw_comparison(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """COMPARISON — structured A vs B."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    left_heading  = content.get("left_heading",  "Option A")
    right_heading = content.get("right_heading", "Option B")
    left_points   = _ensure_list(content.get("left_points",  []))
    right_points  = _ensure_list(content.get("right_points", []))

    # Left
    _rect(slide, 0.3, 1.6, 6.0, 0.7, ACCENT_RED)
    _txt(slide,  left_heading, 0.4, 1.65, 5.8, 0.5,
         size=18, bold=True, color=TEXT_WHITE)
    _rect(slide, 0.3, 2.35, 6.0, 4.75, BG_CARD)
    _bullets(slide, left_points[:6], 0.5, 2.5, 5.7, 4.4, size=16, color=TEXT_LIGHT)

    # VS badge
    _txt(slide, "VS", 6.4, 3.5, 0.5, 0.5,
         size=20, bold=True, color=ACCENT_PURPLE, align=PP_ALIGN.CENTER)

    # Right
    _rect(slide, 7.0, 1.6, 6.0, 0.7, ACCENT_PURPLE)
    _txt(slide,  right_heading, 7.1, 1.65, 5.8, 0.5,
         size=18, bold=True, color=TEXT_WHITE)
    _rect(slide, 7.0, 2.35, 6.0, 4.75, BG_CARD)
    _bullets(slide, right_points[:6], 7.1, 2.5, 5.7, 4.4, size=16, color=TEXT_LIGHT)


def _draw_timeline(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """TIMELINE — horizontal timeline with up to 4 events."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    # Timeline spine
    _rect(slide, 0.5, 3.9, 12.3, 0.06, ACCENT_RED)

    items = content.get("items", [])
    if not items:
        return

    n = min(len(items), 4)
    positions_x = [1.0, 4.0, 7.0, 10.0][:n]

    for i, (item, x) in enumerate(zip(items, positions_x)):
        if isinstance(item, dict):
            date  = str(item.get("date", f"Q{i+1}"))
            event = str(item.get("event", f"Event {i+1}"))
            desc  = str(item.get("description", ""))
        else:
            date, event, desc = str(item), f"Event {i+1}", ""

        # Dot
        _oval(slide, x + 0.85, 3.65, 0.3, 0.3, ACCENT_RED)

        # Date label below dot
        _txt(slide, date, x + 0.5, 4.1, 1.3, 0.4,
             size=14, bold=True, color=ACCENT_RED, align=PP_ALIGN.CENTER)

        # Event card: alternate above/below
        card_text = f"{event}\n{desc}".strip() if desc else event
        if i % 2 == 0:
            _rect(slide, x + 0.2, 2.0, 2.2, 1.5, BG_CARD)
            _txt(slide, card_text, x + 0.3, 2.1, 2.0, 1.2,
                 size=13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)
        else:
            _rect(slide, x + 0.2, 4.7, 2.2, 1.5, BG_CARD)
            _txt(slide, card_text, x + 0.3, 4.8, 2.0, 1.2,
                 size=13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)


def _draw_process(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """PROCESS — step-by-step flow with up to 4 steps."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    steps = content.get("steps", [])
    if not steps:
        return

    n = min(len(steps), 4)
    colors = [ACCENT_RED, ACCENT_PURPLE, ACCENT_BLUE, ACCENT_TEAL]

    for i, step in enumerate(steps[:n]):
        if isinstance(step, dict):
            title = str(step.get("title", f"Step {i+1}"))
            desc  = str(step.get("description", ""))
            num   = str(step.get("step_number", i + 1))
        else:
            title, desc, num = str(step), "", str(i + 1)

        x = 0.4 + i * 3.2
        color = colors[i % len(colors)]

        # Box
        _rect(slide, x, 2.0, 2.8, 4.5, BG_CARD)

        # Number badge
        _oval(slide, x + 1.0, 1.85, 0.8, 0.8, color)
        _txt(slide, num, x + 1.0, 1.9, 0.8, 0.6,
             size=18, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

        # Step title
        _txt(slide, title, x + 0.1, 2.9, 2.6, 0.6,
             size=15, bold=True, color=color, align=PP_ALIGN.CENTER)

        # Description
        if desc:
            _txt(slide, desc, x + 0.1, 3.6, 2.6, 2.6,
                 size=13, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)

        # Arrow between steps
        if i < n - 1:
            _txt(slide, "→", x + 2.85, 3.8, 0.5, 0.6,
                 size=24, bold=True, color=ACCENT_RED, align=PP_ALIGN.CENTER)


def _draw_architecture(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """ARCHITECTURE — layered boxes."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    layers = content.get("layers", [])
    if not layers:
        comps  = content.get("components", [])
        layers = [{"name": c, "description": "", "sub_components": []} for c in comps]
    if not layers:
        return

    row_colors = [ACCENT_RED, ACCENT_PURPLE, ACCENT_BLUE, ACCENT_TEAL]

    for i, layer in enumerate(layers[:4]):
        if isinstance(layer, dict):
            name = str(layer.get("name", f"Layer {i+1}"))
            desc = str(layer.get("description", ""))
            sub  = [str(s) for s in layer.get("sub_components", [])]
        else:
            name, desc, sub = str(layer), "", []

        y = 1.6 + i * 1.3
        color = row_colors[i % len(row_colors)]

        _rect(slide, 0.4, y, 12.5, 1.1, color)
        _txt(slide, name, 0.5, y + 0.05, 4.0, 0.5,
             size=16, bold=True, color=TEXT_WHITE)

        detail = desc
        if sub:
            detail = (detail + " — " if detail else "") + ", ".join(sub[:4])
        if detail:
            _txt(slide, detail, 4.8, y + 0.05, 7.8, 0.5,
                 size=14, color=TEXT_LIGHT)

        if i < min(len(layers), 4) - 1:
            _txt(slide, "↕", 6.4, y + 1.1, 0.5, 0.2,
                 size=12, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def _draw_statistics(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """STATISTICS — KPI metric cards."""
    _bg(slide, BG_DARK)
    _header_bar(slide, slide_def.title)

    statistics = content.get("statistics", [])
    if not statistics:
        statistics = [{"value": b, "label": "", "context": ""}
                      for b in _ensure_list(content.get("bullets", []))]
    if not statistics:
        return

    stat_colors = [ACCENT_RED, ACCENT_PURPLE, ACCENT_GREEN, ACCENT_TEAL]

    for i, stat in enumerate(statistics[:4]):
        if isinstance(stat, dict):
            value   = str(stat.get("value", ""))
            label   = str(stat.get("label", ""))
            context = str(stat.get("context", ""))
        else:
            value, label, context = str(stat), "", ""

        x = 0.4 + i * 3.2
        color = stat_colors[i % len(stat_colors)]

        # Card
        _rect(slide, x, 1.8, 2.9, 4.8, BG_CARD)
        _rect(slide, x, 1.8, 2.9, 0.08, color)

        # Value
        _txt(slide, value, x + 0.1, 2.5, 2.7, 1.4,
             size=48, bold=True, color=color, align=PP_ALIGN.CENTER)

        # Label
        label_text = f"{label}\n{context}".strip() if context else label
        if label_text:
            _txt(slide, label_text, x + 0.1, 4.1, 2.7, 1.8,
                 size=15, color=TEXT_LIGHT, align=PP_ALIGN.CENTER)

    # Context note
    ctx = content.get("context_paragraph", "")
    if ctx:
        _txt(slide, ctx, 0.4, 6.8, 12.5, 0.4,
             size=11, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def _draw_conclusion(slide, content: Dict, slide_def: SlideDefinition) -> None:
    """CONCLUSION — summary + recommendations + next steps + CTA."""
    _bg(slide, BG_DARK)

    # Header
    _rect(slide, 0, 0, 13.333, 1.5, ACCENT_BLUE)
    _rect(slide, 0, 0, 0.1, 1.5, ACCENT_RED)
    _txt(slide, slide_def.title,
         0.3, 0.1, 12.5, 1.1,
         size=28, bold=True, color=TEXT_WHITE)

    summary_points  = _ensure_list(content.get("summary_points", []))
    recommendations = _ensure_list(content.get("recommendations", []))
    next_steps      = _ensure_list(content.get("next_steps", []))
    cta             = content.get("call_to_action", "")

    has_recs = bool(recommendations)

    if has_recs:
        # Three-panel layout: Summary | Recommendations | Next Steps
        # Summary panel (left)
        _rect(slide, 0.3, 1.7, 4.1, 5.1, BG_CARD)
        _txt(slide, "SUMMARY", 0.4, 1.75, 3.9, 0.4,
             size=11, bold=True, color=ACCENT_TEAL)
        if summary_points:
            _bullets(slide, summary_points[:5], 0.4, 2.2, 3.9, 4.3,
                     size=14, color=TEXT_LIGHT)

        # Recommendations panel (centre)
        _rect(slide, 4.6, 1.7, 4.1, 5.1, ACCENT_PURPLE)
        _txt(slide, "RECOMMENDATIONS", 4.7, 1.75, 3.9, 0.4,
             size=11, bold=True, color=TEXT_WHITE)
        if recommendations:
            recs_text = "\n".join(
                f"{i+1}. {r}" for i, r in enumerate(recommendations[:5])
            )
            _txt(slide, recs_text, 4.7, 2.2, 3.9, 4.3,
                 size=14, color=TEXT_WHITE)

        # Next Steps panel (right)
        _rect(slide, 8.9, 1.7, 4.1, 5.1, ACCENT_BLUE)
        _txt(slide, "NEXT STEPS", 9.0, 1.75, 3.8, 0.4,
             size=11, bold=True, color=TEXT_WHITE)
        if next_steps:
            steps_text = "\n".join(
                f"{i+1}. {s}" for i, s in enumerate(next_steps[:5])
            )
            _txt(slide, steps_text, 9.0, 2.2, 3.8, 4.3,
                 size=14, color=TEXT_WHITE)
    else:
        # Two-panel layout (no recommendations): Summary | Next Steps
        _rect(slide, 0.3, 1.7, 7.8, 5.1, BG_CARD)
        _txt(slide, "SUMMARY", 0.4, 1.75, 3.0, 0.4,
             size=11, bold=True, color=ACCENT_RED)
        if summary_points:
            _bullets(slide, summary_points[:6], 0.4, 2.2, 7.5, 4.2,
                     size=16, color=TEXT_LIGHT)

        _rect(slide, 8.4, 1.7, 4.6, 5.1, ACCENT_PURPLE)
        _txt(slide, "NEXT STEPS", 8.5, 1.75, 4.3, 0.4,
             size=11, bold=True, color=TEXT_WHITE)
        if next_steps:
            steps_text = "\n".join(
                f"{i+1}. {s}" for i, s in enumerate(next_steps[:5])
            )
            _txt(slide, steps_text, 8.5, 2.2, 4.3, 3.8, size=15, color=TEXT_WHITE)

    # CTA bar
    _rect(slide, 3.0, 6.95, 7.3, 0.35, ACCENT_RED)
    cta_text = cta if cta else "Thank You — Questions Welcome"
    _txt(slide, cta_text, 3.0, 6.95, 7.3, 0.35,
         size=13, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Dispatch table
# ---------------------------------------------------------------------------

_DRAWERS = {
    # ---- Universal visual archetypes (new — LLM-facing) ----
    "TITLE":             _draw_title,
    "EXECUTIVE_SUMMARY": _draw_executive_summary,   # NEW — slide 2
    "AGENDA":            _draw_agenda,              # NEW — slide 3
    "BULLETS":           _draw_title_content,       # replaces TITLE_AND_CONTENT
    "SECTION_HEADER":    _draw_section_header,
    "CARDS_2_COL":       _draw_two_column,          # replaces TWO_COLUMN / COMPARISON
    "CARDS_3_COL":       _draw_process,             # replaces PROCESS / THREE_COLUMN
    "TABLE":             _draw_architecture,        # closest structured visual
    "STATS":             _draw_statistics,          # replaces STATISTICS
    "IMAGE_TEXT":        _draw_image_text,          # replaces ARCHITECTURE
    "TIMELINE":          _draw_timeline,
    "CONCLUSION":        _draw_conclusion,
    # ---- Legacy names (backward compatibility) ----
    "TITLE_AND_CONTENT": _draw_title_content,
    "TWO_COLUMN":        _draw_two_column,
    "COMPARISON":        _draw_comparison,
    "PROCESS":           _draw_process,
    "ARCHITECTURE":      _draw_architecture,
    "STATISTICS":        _draw_statistics,
}

# Universal archetype aliases for the dark navy builder
_UNIVERSAL_ALIASES = {
    "TITLE_SLIDE": "TITLE", "COVER": "TITLE", "OPENING": "TITLE",
    # Executive Summary aliases
    "EXEC_SUMMARY": "EXECUTIVE_SUMMARY", "KEY_HIGHLIGHTS": "EXECUTIVE_SUMMARY",
    "EXECUTIVE": "EXECUTIVE_SUMMARY", "HIGHLIGHTS": "EXECUTIVE_SUMMARY",
    "OVERVIEW": "EXECUTIVE_SUMMARY",
    # Agenda aliases
    "TOC": "AGENDA", "TABLE_OF_CONTENTS": "AGENDA", "CONTENTS": "AGENDA",
    "TOPICS": "AGENDA",
    # Bullets aliases
    "BULLET_POINTS": "BULLETS", "CONTENT": "BULLETS",
    "GENERAL": "BULLETS", "DEFAULT": "BULLETS", "TITLE_AND_CONTENT": "BULLETS",
    "SPLIT": "CARDS_2_COL", "COLUMNS": "CARDS_2_COL",
    "TWO_COLUMN": "CARDS_2_COL", "COMPARISON": "CARDS_2_COL",
    "THREE_COLUMN": "CARDS_3_COL", "STEPS": "CARDS_3_COL",
    "WORKFLOW": "CARDS_3_COL", "FLOWCHART": "CARDS_3_COL", "PROCESS": "CARDS_3_COL",
    "HEADER": "SECTION_HEADER", "DIVIDER": "SECTION_HEADER", "CHAPTER": "SECTION_HEADER",
    "STATISTICS": "STATS", "METRICS": "STATS", "KPIS": "STATS",
    "NUMBERS": "STATS", "DATA": "STATS",
    "SYSTEM": "IMAGE_TEXT", "TECH_STACK": "IMAGE_TEXT",
    "ARCHITECTURE": "IMAGE_TEXT", "IMAGE": "IMAGE_TEXT", "VISUAL": "IMAGE_TEXT",
    "ROADMAP": "TIMELINE", "MILESTONES": "TIMELINE", "AWARDS": "TIMELINE",
    "GRID": "TABLE", "MATRIX": "TABLE",
    "SUMMARY": "CONCLUSION", "CLOSING": "CONCLUSION", "TAKEAWAYS": "CONCLUSION",
    "RECOMMENDATIONS": "CONCLUSION",
}


# ---------------------------------------------------------------------------
# Main builder
# ---------------------------------------------------------------------------

class PptxBuilderError(Exception):
    pass


def build_presentation(
    plan: PresentationPlan,
    layout_manager: LayoutManager,
) -> bytes:
    """
    Build the final editable .pptx from a PresentationPlan.

    Each slide is drawn FROM SCRATCH using python-pptx shape primitives
    so that LLM-generated content always persists after save/reload.

    Returns raw bytes of the generated .pptx file.
    """
    logger.info("Building '%s' (%d slides)", plan.title, len(plan.slides))

    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    blank_layout = prs.slide_layouts[6]  # Blank

    for slide_def in plan.slides:
        slide = prs.slides.add_slide(blank_layout)

        # Resolve universal archetype key first, then fall back to layout_manager
        raw_lt = slide_def.layout_type.upper().strip()
        layout_type = _UNIVERSAL_ALIASES.get(raw_lt, raw_lt)
        if layout_type not in _DRAWERS:
            # Still unknown — try layout_manager normalisation
            layout_type = layout_manager.normalise_layout_type(slide_def.layout_type)
        draw_fn = _DRAWERS.get(layout_type, _draw_title_content)
        content  = slide_def.content or {}

        logger.debug(
            "Slide %d [%s]: %s | keys=%s",
            slide_def.slide_number, layout_type,
            slide_def.title, list(content.keys())
        )

        try:
            draw_fn(slide, content, slide_def)
        except Exception as e:
            logger.warning(
                "Slide %d (%s) draw failed (%s: %s) — falling back to title+content",
                slide_def.slide_number, slide_def.title, type(e).__name__, e
            )
            try:
                _draw_title_content(slide, content, slide_def)
            except Exception as e2:
                logger.error("Fallback also failed for slide %d: %s", slide_def.slide_number, e2)
                # Last resort: at least put the title
                try:
                    _bg(slide, BG_DARK)
                    _txt(slide, slide_def.title, 1.0, 2.0, 11.0, 2.0,
                         size=32, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)
                except Exception:
                    pass

        # Speaker notes
        if slide_def.speaker_notes:
            try:
                prs.slides[-1].notes_slide.notes_text_frame.text = slide_def.speaker_notes
            except Exception:
                pass

    try:
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        data = buf.read()
        logger.info("PPTX built: %d bytes, %d slides", len(data), len(prs.slides))
        return data
    except Exception as e:
        raise PptxBuilderError(f"Failed to save presentation: {e}") from e
