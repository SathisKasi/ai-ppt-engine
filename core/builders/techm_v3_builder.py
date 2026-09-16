"""
core/builders/techm_v3_builder.py — Dynamic slide builder for TechM_RefPPT-V3.

Architecture:
  - Loads TechM_RefPPT-V3.pptx containing 3 master slides:
      Slide 1 (idx 0): Cover Page
      Slide 2 (idx 1): 1_Title and Content (Content Canvas Base)
      Slide 3 (idx 2): Thankyou
  - OpenXML Deep-Cloning:
      1. Clones Slide 1 for Cover Page -> updates Title and Date in-place.
      2. Duplicates Slide 2 N times for N content slides -> updates header title,
         then invokes DynamicGeometryEngine to render custom visual layouts
         (KPI ribbons, multi-column cards, process pipelines, comparison splits, tables).
      3. Clones Slide 3 for the Thank You closing slide -> 100% intact corporate branding.
  - Prunes original 3 template slides from OpenXML package.
  - Returns binary bytes via io.BytesIO().
"""
from __future__ import annotations

import copy
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

from pptx import Presentation
from pptx.oxml.ns import qn

import config
from core.renderers.dynamic_geometry_engine import DynamicGeometryEngine
from llm.dynamic_layout_schemas import (
    CardItem,
    ComparisonColumn,
    DynamicLayoutComposition,
    DynamicPresentationPlan,
    DynamicSlideDefinition,
    HeroBlock,
    KPIMetricItem,
    ProcessStepItem,
    TableDataGrid,
    TimelineItem,
)
from llm.schemas import PresentationPlan, SlideDefinition
from utils.logging_utils import get_logger

logger = get_logger(__name__)


def _clone_slide(prs: Presentation, source_idx: int) -> Any:
    """
    Deep-clones a slide from the presentation catalog at source_idx.
    Clones media relationships and remaps rId references to prevent broken images.
    """
    source_slide = prs.slides[source_idx]
    slide_layout = source_slide.slide_layout
    new_slide = prs.slides.add_slide(slide_layout)

    # Map old relationship rId to new slide relationship rId
    rId_map = {}
    for rId, rel in source_slide.part.rels.items():
        if (
            "slideLayout" not in rel.target_ref
            and "notesSlide" not in rel.target_ref
            and "notesMaster" not in rel.target_ref
        ):
            try:
                new_rId = new_slide.part.relate_to(rel.target_part, rel.reltype)
                rId_map[rId] = new_rId
            except Exception as e:
                logger.debug("Could not relate part for rId %s: %s", rId, e)

    # Clear default placeholders injected by add_slide
    for s in list(new_slide.shapes):
        sp = s.element
        sp.getparent().remove(sp)

    # Deep-copy all shapes and remap relationship attributes
    for shape in source_slide.shapes:
        new_el = copy.deepcopy(shape.element)
        for elem in new_el.xpath(".//*[@r:embed or @r:link or @r:id]"):
            for attr in list(elem.attrib.keys()):
                if attr.endswith("embed") or attr.endswith("link") or attr.endswith("}id") or attr == "r:id":
                    old_val = elem.attrib[attr]
                    if old_val in rId_map:
                        elem.attrib[attr] = rId_map[old_val]
        new_slide.shapes._spTree.append(new_el)

    return new_slide


def _set_tf_text(shape: Any, text: str) -> bool:
    """Sets text in-place preserving existing font styles."""
    if not hasattr(shape, "text_frame"):
        return False
    tf = shape.text_frame
    text_str = str(text).strip()
    if not tf.paragraphs:
        p = tf.add_paragraph()
        run = p.add_run()
        run.text = text_str
        return True

    p0 = tf.paragraphs[0]
    if p0.runs:
        p0.runs[0].text = text_str
        for r in p0.runs[1:]:
            r.text = ""
    else:
        p0.text = text_str

    for extra_p in tf.paragraphs[1:]:
        for r in extra_p.runs:
            r.text = ""
        extra_p.text = ""
    return True


def _convert_legacy_slide_to_dynamic(slide_def: SlideDefinition) -> DynamicSlideDefinition:
    """Converts a standard SlideDefinition into a DynamicSlideDefinition."""
    lt = (slide_def.layout_type or "BULLETS").upper()
    content = slide_def.content or {}

    if lt in ("STATS", "STATISTICS"):
        stats = content.get("statistics", [])
        metrics = [
            KPIMetricItem(
                value=str(s.get("value", "")),
                label=str(s.get("label", "")),
                delta=str(s.get("context", "")) if s.get("context") else None,
            )
            for s in stats
            if isinstance(s, dict)
        ]
        comp = DynamicLayoutComposition(
            pattern="METRIC_RIBBON_AND_CARDS",
            metrics=metrics,
            cards=[
                CardItem(
                    title="Key Performance Takeaway",
                    bullets=[slide_def.purpose or "Key quantified benchmark metrics."],
                )
            ],
        )
    elif lt in ("CARDS_2_COL", "TWO_COLUMN", "COMPARISON"):
        left_h = content.get("left_heading", "Option A")
        right_h = content.get("right_heading", "Option B")
        left_b = content.get("left_bullets", [])
        right_b = content.get("right_bullets", [])
        comp = DynamicLayoutComposition(
            pattern="COMPARISON_SPLIT",
            comparison=[
                ComparisonColumn(heading=left_h, tag="Option A", bullets=left_b),
                ComparisonColumn(heading=right_h, tag="Option B", bullets=right_b),
            ],
        )
    elif lt in ("CARDS_3_COL", "THREE_COLUMN", "PROCESS"):
        steps_data = content.get("steps", [])
        steps = []
        for i, st in enumerate(steps_data):
            if isinstance(st, dict):
                steps.append(
                    ProcessStepItem(
                        step_number=str(st.get("step_number", i + 1)),
                        title=str(st.get("title", f"Step {i + 1}")),
                        description=str(st.get("description", "")),
                    )
                )
        if not steps:
            # Check for generic components or bullets
            bullets = content.get("bullets", [])
            chunk = max(1, len(bullets) // 3)
            cards = [
                CardItem(title="Pillar 1", bullets=bullets[:chunk]),
                CardItem(title="Pillar 2", bullets=bullets[chunk : 2 * chunk]),
                CardItem(title="Pillar 3", bullets=bullets[2 * chunk :]),
            ]
            comp = DynamicLayoutComposition(pattern="MULTI_COLUMN_CARDS", cards=cards)
        else:
            comp = DynamicLayoutComposition(pattern="PROCESS_PIPELINE", steps=steps)
    elif lt == "TABLE":
        t_data = content.get("table_data", {})
        headers = t_data.get("headers", ["Category", "Details"])
        rows = t_data.get("rows", [["Item 1", "Value 1"]])
        comp = DynamicLayoutComposition(
            pattern="DATA_MATRIX_TABLE",
            table=TableDataGrid(headers=headers, rows=rows),
        )
    elif lt in ("TIMELINE", "AWARDS"):
        items_data = content.get("items", [])
        timeline = []
        for it in items_data:
            if isinstance(it, dict):
                timeline.append(
                    TimelineItem(
                        date=str(it.get("date", "")),
                        event=str(it.get("event", "Milestone")),
                        description=str(it.get("description", "")),
                    )
                )
        comp = DynamicLayoutComposition(pattern="TIMELINE_ROADMAP", timeline=timeline)
    elif lt == "EXECUTIVE_SUMMARY":
        highlights = content.get("highlights", [])
        purpose = content.get("purpose_statement", slide_def.purpose or "")
        key_metric = content.get("key_metric", "")
        metrics = [KPIMetricItem(value=key_metric, label="Target Metric")] if key_metric else []
        cards = [
            CardItem(title="Executive Highlights", bullets=highlights),
            CardItem(title="Strategic Scope", bullets=[purpose] if purpose else []),
        ]
        comp = DynamicLayoutComposition(
            pattern="METRIC_RIBBON_AND_CARDS" if metrics else "MULTI_COLUMN_CARDS",
            metrics=metrics if metrics else None,
            cards=cards,
        )
    else:
        # Default bullets
        bullets = content.get("bullets", [])
        comp = DynamicLayoutComposition(
            pattern="MULTI_COLUMN_CARDS",
            cards=[
                CardItem(
                    title="Key Strategic Points",
                    bullets=bullets if bullets else [slide_def.purpose or "Core topic overview"],
                )
            ],
        )

    return DynamicSlideDefinition(
        slide_number=slide_def.slide_number,
        title=slide_def.title,
        executive_takeaway=slide_def.purpose,
        layout_pattern=comp.pattern,
        composition=comp,
        speaker_notes=slide_def.speaker_notes,
    )


class TechMV3Builder:
    """
    100% Fidelity Slide Cloner & Dynamic Geometry Mutator for TechM_RefPPT-V3.pptx.
    """

    def __init__(self, template_path: Optional[Path] = None):
        self.template_path = template_path or config.TECHM_V3_TEMPLATE_FILE
        if not self.template_path.exists():
            raise FileNotFoundError(f"TechM V3 Template not found at: {self.template_path}")
        self.geometry_engine = DynamicGeometryEngine()

    def build(self, plan: Union[DynamicPresentationPlan, PresentationPlan]) -> bytes:
        """
        Builds the presentation by deep-cloning TechM_RefPPT-V3 slides and rendering dynamic layouts.
        """
        logger.info("Opening TechM V3 master template: %s", self.template_path)
        prs = Presentation(str(self.template_path))
        initial_slide_count = len(prs.slides)
        logger.info("Template catalog slides loaded: %d", initial_slide_count)

        if initial_slide_count < 3:
            raise ValueError(
                f"TechM_RefPPT-V3.pptx must have at least 3 slides (Cover, Content, Thank You). Found: {initial_slide_count}"
            )

        # Normalize plan slides to DynamicSlideDefinition
        if isinstance(plan, DynamicPresentationPlan):
            dynamic_slides = plan.slides
            pres_title = plan.title
            pres_date = plan.date or datetime.now().strftime("%d%b%Y").upper()
        else:
            pres_title = plan.title
            pres_date = datetime.now().strftime("%d%b%Y").upper()
            # Filter out Title or Closing if present in legacy slides
            dynamic_slides = []
            for s in plan.slides:
                lt_up = (s.layout_type or "").upper()
                if lt_up in ("TITLE", "CLOSING"):
                    continue
                dynamic_slides.append(_convert_legacy_slide_to_dynamic(s))

        generated_slides = []

        # ===================================================================
        # 1. Deep-Clone Slide 1 (Cover Page)
        # ===================================================================
        logger.info("[V3 Builder] Cloning Cover Page (Slide 1)")
        cover_slide = _clone_slide(prs, 0)
        generated_slides.append(cover_slide)

        # Mutate Title and Date
        for shape in cover_slide.shapes:
            if shape.name == "Title 1" or (shape.is_placeholder and shape.placeholder_format.idx == 0):
                _set_tf_text(shape, pres_title)
            elif shape.name == "Text Placeholder 2" or "Date" in shape.name or (shape.is_placeholder and shape.placeholder_format.idx == 1):
                _set_tf_text(shape, pres_date)

        # ===================================================================
        # 2. Duplicate Slide 2 (Content Base Canvas) N times
        # ===================================================================
        for s_idx, slide_def in enumerate(dynamic_slides, start=2):
            logger.info(
                "[V3 Builder] Duplicating Slide 2 for content slide #%d: '%s' (Pattern: %s)",
                s_idx,
                slide_def.title[:40],
                slide_def.layout_pattern,
            )
            cloned_slide = _clone_slide(prs, 1)
            generated_slides.append(cloned_slide)

            # Set Slide Title in Text Placeholder 9 (Header Bar)
            title_set = False
            for shape in cloned_slide.shapes:
                if shape.name == "Text Placeholder 9" or (shape.is_placeholder and shape.placeholder_format.idx == 1):
                    _set_tf_text(shape, slide_def.title)
                    title_set = True
                    break
            if not title_set:
                # Fallback: check first placeholder
                for shape in cloned_slide.shapes:
                    if shape.is_placeholder:
                        _set_tf_text(shape, slide_def.title)
                        break

            # Render Dynamic Layout Components inside the safe canvas
            try:
                self.geometry_engine.render_slide_content(cloned_slide, slide_def)
            except Exception as e:
                logger.error("[V3 Builder] Failed rendering dynamic content for slide %d: %s", s_idx, e, exc_info=True)

            # Attach Speaker Notes
            if slide_def.speaker_notes:
                try:
                    cloned_slide.notes_slide.notes_text_frame.text = slide_def.speaker_notes
                except Exception:
                    pass

        # ===================================================================
        # 3. Deep-Clone Slide 3 (Thank You / Closing Slide)
        # ===================================================================
        logger.info("[V3 Builder] Appending Thank You closing slide (Slide 3)")
        closing_slide = _clone_slide(prs, 2)
        generated_slides.append(closing_slide)

        # ===================================================================
        # 4. Prune the original 3 template blueprint slides
        # ===================================================================
        rId_attr = qn("r:id")
        sldIdLst = prs.slides._sldIdLst
        for _ in range(initial_slide_count):
            sldId = sldIdLst[0]
            rId = sldId.get(rId_attr)
            if rId:
                prs.part.drop_rel(rId)
            sldIdLst.remove(sldId)

        logger.info(
            "[V3 Builder] Successfully generated %d slides. Blueprint slides pruned.",
            len(generated_slides),
        )

        # ===================================================================
        # 5. Serialize into in-memory io.BytesIO()
        # ===================================================================
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return buf.read()
