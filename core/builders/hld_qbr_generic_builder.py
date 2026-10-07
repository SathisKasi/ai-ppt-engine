"""
core/builders/hld_qbr_generic_builder.py — Generic, catalog-driven renderer
for the HLD QBR template.

Consumes llm.hld_qbr_generic_schemas.GenericHLDQBRPlan (produced by
core/presentation_planner_hld_qbr_generic.py) + templates/hld_qbr_assets/
inventory/layout_capability_catalog.json. Renders each slide by looking up
its shapes by shape_id (stable across clone, since _clone_slide deep-copies
the exact template XML) and setting values generically — no per-archetype
dispatch functions, no named topic fields.

Scope note (v1): covers the "slot-fill" pattern only — slides whose content
goes into the template's OWN pre-existing sample shapes (text/table/chart).
Blank-canvas composition (ad-hoc extra charts/tables beyond the slide's own
native ones, stat-highlight badges, process-flow chevrons, and the old rich
executive-summary bullet-cards) is not yet wired into this generic path.
"""
from __future__ import annotations

import copy
import io
import re
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

import config
from core.builders.hld_qbr_builder import (
    TEMPLATE_SERIES_COLORS,
    _clone_slide,
    _fill_table_rows,
    _get_ordinal_date,
    _set_first_run_text,
    _strip_all_highlights,
    _strip_decorative_connectors,
    _strip_guidance_shapes,
)
from core.hld_qbr_catalog import STRUCTURAL_ONLY_SLIDE_IDS, get_entry
from llm.hld_qbr_generic_schemas import GenericHLDQBRPlan, SlideAssignment
from utils.logging_utils import get_logger

logger = get_logger(__name__)

# Default font sizes by slot kind, used when a slot's sample run had no
# explicit font.size captured (common — many runs inherit size from the
# layout/master rather than setting it directly on the run).
_DEFAULT_FONT_PT = {
    "placeholder_title": 20,
    "empty_title_box": 20,
    "text_box": 12,
    "auto_shape": 12,
}
TEXT_COLOR = RGBColor(0, 43, 73)


def _shape_by_id(slide: Any, shape_id: int) -> Optional[Any]:
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    return None


def _clear_template_text(shapes: Any) -> None:
    """Blank cloned sample content while preserving the template geometry and
    formatting. Generated values are the only text subsequently written."""
    for shape in list(shapes):
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = ""
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    cell.text = ""
        if getattr(shape, "shape_type", None) == 6:  # GROUP
            _clear_template_text(shape.shapes)


def _remove_unfilled_charts(slide: Any, assignment: Optional[SlideAssignment]) -> None:
    """A cloned chart retains its source workbook labels unless it is replaced
    with source-grounded data, so omit it when the assignment has no chart."""
    has_chart_data = bool(assignment and assignment.chart_categories and assignment.chart_series)
    if has_chart_data:
        return
    for shape in list(slide.shapes):
        if getattr(shape, "has_chart", False):
            shape.element.getparent().remove(shape.element)


def _set_text_generic(shape: Any, text: str, kind: str, sample_max_chars: Optional[int]) -> None:
    """Sets text generically, auto-scaling font size down a step when the new
    text is noticeably longer than what the slot was sized for. Cleans up unused
    template paragraphs so extra blank runs do not break vertical centering."""
    if not getattr(shape, "has_text_frame", False):
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    is_title = kind in ("placeholder_title", "empty_title_box")
    base_pt = _DEFAULT_FONT_PT.get(kind, 12)
    if sample_max_chars and len(text) > sample_max_chars:
        base_pt = max(8, base_pt - 2)

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        lines = [""]

    tf.word_wrap = True
    if kind == "auto_shape":
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Pt(6)
        tf.margin_right = Pt(6)
        tf.margin_top = Pt(4)
        tf.margin_bottom = Pt(4)

    existing_paras = list(tf.paragraphs)

    for i in range(min(len(lines), len(existing_paras))):
        para = existing_paras[i]
        line = lines[i]
        if para.runs:
            existing_size = para.runs[0].font.size
            para.runs[0].text = line
            for r in para.runs[1:]:
                r.text = ""
            if sample_max_chars and len(text) > sample_max_chars:
                cur_pt = existing_size.pt if existing_size else base_pt
                para.runs[0].font.size = Pt(max(8, cur_pt - 2))
            elif existing_size is None and line:
                para.runs[0].font.size = Pt(base_pt)
            if not is_title:
                para.runs[0].font.bold = False
        else:
            run = para.add_run()
            run.text = line
            run.font.size = Pt(base_pt)
            if not is_title:
                run.font.bold = False

    if len(lines) > len(existing_paras):
        sample_run = existing_paras[0].runs[0] if existing_paras and existing_paras[0].runs else None
        for extra_line in lines[len(existing_paras):]:
            p = tf.add_paragraph()
            run = p.add_run()
            run.text = extra_line
            run.font.size = Pt(base_pt)
            if sample_run:
                if sample_run.font.name:
                    run.font.name = sample_run.font.name
                run.font.bold = is_title
            else:
                run.font.bold = is_title
    elif len(existing_paras) > len(lines):
        for extra_p in existing_paras[max(1, len(lines)):]:
            try:
                extra_p._p.getparent().remove(extra_p._p)
            except Exception:
                pass


def _render_simple_slots(slide: Any, entry: Dict[str, Any], assignment: SlideAssignment) -> None:
    seen_texts = set()
    for slot in entry["slots"]:
        if slot["kind"] in ("table", "chart", "placeholder_title", "empty_title_box"):
            continue
        value = assignment.slot_values.get(slot["slot_id"])
        shape = _shape_by_id(slide, slot["shape_id"])
        if shape is None:
            continue
        # Deduplication guard: if an identical non-trivial statement was already rendered on this slide,
        # prune the duplicate shape rather than stamping identical duplicate cards
        norm_val = value.strip().lower() if value else ""
        if value and (norm_val in seen_texts and len(norm_val) > 20):
            try:
                shape.element.getparent().remove(shape.element)
            except Exception:
                pass
            continue
        if value:
            if shape.has_text_frame:
                seen_texts.add(norm_val)
                _set_text_generic(shape, value, slot["kind"], slot.get("max_chars"))
        else:
            # Unfilled optional content shape — remove it rather than leave a
            # visible blank box occupying layout space.
            try:
                shape.element.getparent().remove(shape.element)
            except Exception:
                pass


def _render_auto_numbers(slide: Any, entry: Dict[str, Any], filled_count: int) -> None:
    """Sequence badges (e.g. achievement milestone numbers) next to whichever
    real content slots got filled — purely positional, auto-incrementing."""
    auto_ids = entry.get("auto_number_shape_ids") or []
    for i, shape_id in enumerate(auto_ids):
        shape = _shape_by_id(slide, shape_id)
        if shape is None or not shape.has_text_frame:
            continue
        if i < filled_count:
            _set_first_run_text(shape, str(i + 1))
        else:
            try:
                shape.element.getparent().remove(shape.element)
            except Exception:
                pass


def _render_repeat_groups(
    slide: Any,
    entry: Dict[str, Any],
    assignment: SlideAssignment,
    group_text_children_map: Optional[Dict[int, List[int]]] = None,
) -> None:
    group_map = group_text_children_map or {}
    for rg in entry["repeat_groups"]:
        member_ids = rg["member_shape_ids"]
        # Deduplicate repeat items: don't render identical duplicate cards
        unique_items = []
        seen_card_signatures = set()
        for it in assignment.repeat_items:
            vals = [str(v).strip().lower() for v in it.values() if str(v).strip()]
            sig = tuple(vals)
            if sig and sig in seen_card_signatures:
                continue
            if sig:
                seen_card_signatures.add(sig)
                unique_items.append(it)
        items = unique_items[: len(member_ids)]
        for i, shape_id in enumerate(member_ids):
            member = _shape_by_id(slide, shape_id)
            if member is None:
                continue
            if i >= len(items):
                try:
                    member.element.getparent().remove(member.element)
                except Exception:
                    pass
                continue
            item_values = items[i]
            target_child_ids = group_map.get(shape_id)
            if target_child_ids:
                child_by_id = {c.shape_id: c for c in member.shapes}
                text_children = [child_by_id[cid] for cid in target_child_ids if cid in child_by_id]
            else:
                # Fallback: filter out non-text shapes (Oval icons, Freeform graphics, lines)
                candidate_children = [
                    s for s in member.shapes
                    if getattr(s, "has_text_frame", False)
                    and getattr(s, "shape_type", None) not in (5, 9)
                    and not getattr(s, "name", "").lower().startswith("oval")
                    and not getattr(s, "name", "").lower().startswith("freeform")
                ]
                text_children = sorted(
                    candidate_children,
                    key=lambda s: (round(s.top / 914400, 2), s.left),
                )
            for slot_idx, (slot_def, child) in enumerate(zip(rg["item_slots"], text_children)):
                value = item_values.get(slot_def["slot_id"])
                if not value and slot_idx < len(item_values):
                    # Positional fallback if LLM returned descriptive keys (e.g. "title", "text")
                    candidate_val = list(item_values.values())[slot_idx]
                    if candidate_val:
                        value = str(candidate_val)
                if value:
                    is_title = (slot_def["slot_id"] == "item_slot_1" or slot_idx == 0)
                    kind = "placeholder_title" if is_title else "auto_shape"
                    _set_text_generic(child, value, kind, slot_def.get("max_chars"))
                    if is_title and getattr(child, "has_text_frame", False):
                        if child.text_frame.paragraphs and child.text_frame.paragraphs[0].runs:
                            child.text_frame.paragraphs[0].runs[0].font.bold = True


def _render_table(slide: Any, entry: Dict[str, Any], assignment: SlideAssignment) -> None:
    table_slot = next((s for s in entry["slots"] if s["kind"] == "table"), None)
    rendered_shape_id = None
    if table_slot is not None and assignment.table_rows:
        shape = _shape_by_id(slide, table_slot["shape_id"])
        if shape is not None and shape.has_table:
            rendered_shape_id = shape.shape_id
            tbl = shape.table
            orig_rows = len(tbl.rows)
            orig_height = shape.height
            headers = assignment.table_headers or []
            for c_idx, header in enumerate(headers):
                if c_idx < len(tbl.columns):
                    cell = tbl.cell(0, c_idx)
                    cell.text = str(header)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            r.font.name = "Verdana"
                            r.font.size = Pt(11.0)
                            r.font.bold = True
                            r.font.color.rgb = TEXT_COLOR
            _fill_table_rows(tbl, assignment.table_rows, start_row=1, prune_unused_rows=True, text_color=TEXT_COLOR)
            if orig_rows > 0 and len(tbl.rows) < orig_rows:
                shape.height = min(orig_height, int(orig_height * (len(tbl.rows) / orig_rows)))

    # Remove any unpopulated / ghost table shapes on this slide so empty grids never appear
    for shape in list(slide.shapes):
        if getattr(shape, "has_table", False) and shape.shape_id != rendered_shape_id:
            try:
                shape.element.getparent().remove(shape.element)
            except Exception:
                pass


def _render_chart(slide: Any, entry: Dict[str, Any], assignment: SlideAssignment) -> None:
    chart_slot = next((s for s in entry["slots"] if s["kind"] == "chart"), None)
    if chart_slot is None or not (assignment.chart_categories and assignment.chart_series):
        return
    shape = _shape_by_id(slide, chart_slot["shape_id"])
    if shape is None or not shape.has_chart:
        return
    chart = shape.chart
    cd = CategoryChartData()
    cd.categories = assignment.chart_categories
    for series in assignment.chart_series:
        cd.add_series(series.name, series.values)
    chart.replace_data(cd)
    chart.has_title = False
    if len(assignment.chart_series) > 1:
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
    try:
        if chart.plots:
            for s_idx, series in enumerate(chart.plots[0].series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = TEMPLATE_SERIES_COLORS[s_idx % len(TEMPLATE_SERIES_COLORS)]
    except Exception:
        pass


def _render_content_slide(prs: Presentation, assignment: SlideAssignment, slide_number: Optional[int] = None) -> bool:
    entry = get_entry(assignment.slide_id)
    if entry is None:
        logger.warning("Skipping unknown slide_id %s", assignment.slide_id)
        return False
    slide = _clone_slide(prs, entry["source_slide_index"])
    _strip_guidance_shapes(slide)
    _strip_decorative_connectors(slide)

    # Capture group text-bearing children BEFORE clearing text so decorative icons/freeforms are ignored
    group_text_children_map: Dict[int, List[int]] = {}
    for shape in slide.shapes:
        if getattr(shape, "shape_type", None) == 6:  # GROUP
            valid_children = [
                c for c in shape.shapes
                if getattr(c, "has_text_frame", False) and c.text_frame.text.strip()
            ]
            valid_children.sort(key=lambda c: (round(c.top / 914400, 2), c.left))
            group_text_children_map[shape.shape_id] = [c.shape_id for c in valid_children]

    _clear_template_text(slide.shapes)
    _remove_unfilled_charts(slide, assignment)

    # Set sequential slide number if provided (must happen after clearing template text)
    if slide_number is not None:
        for shape in slide.shapes:
            if "slide number" in shape.name.lower() or (
                getattr(shape, "is_placeholder", False)
                and getattr(shape.placeholder_format, "type", None) in (13, 16)
            ):
                if shape.has_text_frame:
                    _set_first_run_text(shape, str(slide_number))

    if assignment.title is not None:
        title_slot = next(
            (s for s in entry["slots"] if s["kind"] in ("placeholder_title", "empty_title_box")), None
        )
        if title_slot is not None:
            shape = _shape_by_id(slide, title_slot["shape_id"])
            if shape is not None and shape.has_text_frame:
                _set_text_generic(shape, assignment.title.upper(), title_slot["kind"], title_slot.get("max_chars"))

    _render_simple_slots(slide, entry, assignment)
    filled_simple_count = sum(1 for v in assignment.slot_values.values() if v)
    _render_auto_numbers(slide, entry, filled_simple_count)
    _render_repeat_groups(slide, entry, assignment, group_text_children_map)
    _render_table(slide, entry, assignment)
    _render_chart(slide, entry, assignment)
    return True


def _render_cover(prs: Presentation, plan: GenericHLDQBRPlan) -> None:
    entry = get_entry("slide_00")
    slide = _clone_slide(prs, entry["source_slide_index"])
    _clear_template_text(slide.shapes)
    date_str = plan.date.strip() if plan.date else _get_ordinal_date()
    for slot in entry["slots"]:
        shape = _shape_by_id(slide, slot["shape_id"])
        if shape is None or not shape.has_text_frame:
            continue
        if "date" in slot["slot_id"]:
            _set_first_run_text(shape, date_str)
        elif "title_5" in slot["slot_id"]:
            # Title 5 overlaps the primary hero title (Rectangle 6).
            # Use it for facility/subtitle if present; otherwise remove it to avoid duplicate text.
            if plan.facility_name:
                _set_text_generic(shape, plan.facility_name, slot["kind"], slot.get("max_chars"))
            else:
                try:
                    shape.element.getparent().remove(shape.element)
                except Exception:
                    pass
        else:
            _set_text_generic(shape, plan.presentation_title, slot["kind"], slot.get("max_chars"))


def _render_agenda(prs: Presentation, agenda_title: str, agenda_items: List[str]) -> None:
    entry = get_entry("slide_01")
    slide = _clone_slide(prs, entry["source_slide_index"])
    _clear_template_text(slide.shapes)

    # Update slide number placeholder to 2
    for shape in slide.shapes:
        if "slide number" in shape.name.lower() or (
            getattr(shape, "is_placeholder", False)
            and getattr(shape.placeholder_format, "type", None) in (13, 16)
        ):
            if shape.has_text_frame:
                _set_first_run_text(shape, "2")

    topics_slot = max(entry["slots"], key=lambda s: s.get("max_chars") or 0)
    title_slot = next((slot for slot in entry["slots"] if slot != topics_slot), None)
    if title_slot is not None:
        title_shape = _shape_by_id(slide, title_slot["shape_id"])
        if title_shape is not None and title_shape.has_text_frame:
            display_title = "AGENDA" if not agenda_title or agenda_title.lower() in ("agenda", "today's discussion") else agenda_title.upper()
            _set_text_generic(title_shape, display_title, "placeholder_title", 40)
    shape = _shape_by_id(slide, topics_slot["shape_id"])
    if shape is not None and shape.has_text_frame and agenda_items:
        tf = shape.text_frame
        # Clean and deduplicate agenda items
        cleaned_lines: List[str] = []
        seen = set()
        for t in agenda_items:
            s = re.sub(r"^\d+[\.\)]\s*", "", t.strip())
            if s and s.lower() not in seen:
                seen.add(s.lower())
                cleaned_lines.append(s)

        # Enforce maximum 6 items to protect visual breathing room and prevent footer collisions
        lines = cleaned_lines[:6]
        if not lines:
            return

        # Extract template's master bullet paragraph XML properties to replicate across all paragraphs
        sample_pPr = None
        if tf.paragraphs and tf.paragraphs[0]._p.find(qn("a:pPr")) is not None:
            sample_pPr = copy.deepcopy(tf.paragraphs[0]._p.find(qn("a:pPr")))

        # Adaptive typography & spacing based on item count
        n = len(lines)
        if n <= 3:
            ag_sz = Pt(17)
            line_spc = 1.8
            spc_after = Pt(14)
        elif n <= 4:
            ag_sz = Pt(16)
            line_spc = 1.7
            spc_after = Pt(10)
        elif n <= 5:
            ag_sz = Pt(14)
            line_spc = 1.5
            spc_after = Pt(8)
        else:  # 6 items
            ag_sz = Pt(13)
            line_spc = 1.35
            spc_after = Pt(6)

        existing_paras = list(tf.paragraphs)

        # 1. Populate matching existing paragraphs cleanly
        for i in range(min(len(lines), len(existing_paras))):
            para = existing_paras[i]
            if sample_pPr is not None:
                pPr = para._p.find(qn("a:pPr"))
                if pPr is not None:
                    para._p.remove(pPr)
                para._p.insert(0, copy.deepcopy(sample_pPr))
            if para.runs:
                para.runs[0].text = lines[i]
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run().text = lines[i]
            para.line_spacing = line_spc
            para.space_after = spc_after
            for r in para.runs:
                r.font.name = "Verdana"
                r.font.size = ag_sz
                r.font.color.rgb = RGBColor(0, 43, 73)

        # 2. Add extra paragraphs if more lines than template
        if len(lines) > len(existing_paras):
            for extra_line in lines[len(existing_paras):]:
                p = tf.add_paragraph()
                if sample_pPr is not None:
                    pPr = p._p.find(qn("a:pPr"))
                    if pPr is not None:
                        p._p.remove(pPr)
                    p._p.insert(0, copy.deepcopy(sample_pPr))
                p.line_spacing = line_spc
                p.space_after = spc_after
                run = p.add_run()
                run.text = extra_line
                run.font.name = "Verdana"
                run.font.size = ag_sz
                run.font.color.rgb = RGBColor(0, 43, 73)

        # 3. Remove excess template paragraphs so no blank gaps or ghost bullets exist
        elif len(existing_paras) > len(lines):
            for extra_p in existing_paras[len(lines):]:
                try:
                    extra_p._p.getparent().remove(extra_p._p)
                except Exception:
                    pass


# Brand tagline + legal footer on the closing slide — required verbatim
# copy per guardrails.md, not per-engagement sample content, so it is never
# routed through the LLM like everything else on a generated slide.
_CLOSING_TAGLINE = "MOVING OUR WORLD FORWARD BY DELIVERING WHAT MATTERS\u2122"
_CLOSING_LEGAL_FOOTER = (
    "Proprietary and Confidential: This presentation may not be used or disclosed to "
    "other than employees or customers, unless expressly authorized by UPS.\u000b"
    "\u00a9 2026 United Parcel Service of America, Inc. UPS, the UPS brandmark and the "
    "color brown are trademarks of United Parcel Service of America, Inc. All rights reserved."
)


def _render_closing(prs: Presentation) -> None:
    entry = get_entry("slide_30")
    slide = _clone_slide(prs, entry["source_slide_index"])
    _clear_template_text(slide.shapes)
    texts_by_max_chars_desc = sorted(entry["slots"], key=lambda s: -(s.get("max_chars") or 0))
    for slot, text in zip(texts_by_max_chars_desc, (_CLOSING_LEGAL_FOOTER, _CLOSING_TAGLINE)):
        shape = _shape_by_id(slide, slot["shape_id"])
        if shape is not None and shape.has_text_frame:
            _set_text_generic(shape, text, slot["kind"], slot.get("max_chars"))


class HLDQBRGenericBuilder:
    """Generic counterpart to core.builders.hld_qbr_builder.HLDQBRBuilder."""

    def build(self, plan: GenericHLDQBRPlan) -> bytes:
        prs = Presentation(str(config.HLD_QBR_TEMPLATE_FILE))
        initial_count = len(prs.slides)

        content_slides = [s for s in plan.slides if s.slide_id not in STRUCTURAL_ONLY_SLIDE_IDS]

        _render_cover(prs, plan)
        agenda_items = plan.agenda_topics if plan.agenda_topics else [s.title for s in content_slides if s.title]
        _render_agenda(prs, "AGENDA", agenda_items)

        slide_idx = 3
        for assignment in content_slides:
            if _render_content_slide(prs, assignment, slide_number=slide_idx):
                slide_idx += 1
        _render_closing(prs)

        # Prune the original 60 source slides, leaving only the generated ones
        # (new slides are always appended, so the originals are the first
        # initial_count entries in the slide id list).
        rId_attr = qn("r:id")
        sldIdLst = prs.slides._sldIdLst
        for _ in range(initial_count):
            sldId = sldIdLst[0]
            rId = sldId.get(rId_attr)
            if rId:
                prs.part.drop_rel(rId)
            sldIdLst.remove(sldId)

        logger.info("[HLD QBR Generic Builder] Built %d slides.", len(prs.slides))

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
