"""
core/document_parser.py — Extract text from uploaded documents.

Supports:
- .txt  (with encoding detection)
- .docx (python-docx)
- .pdf  (PyMuPDF primary, pdfplumber fallback)
- .pptx (python-pptx text extraction)
- .xlsx (openpyxl worksheet and cell extraction)

All parsers return a clean, normalized plain-text string.
Raises DocumentParseError on unrecoverable failures.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from utils.logging_utils import get_logger
from utils.text_utils import clean_text

logger = get_logger(__name__)


class DocumentParseError(Exception):
    """Raised when document parsing fails."""


# ---------------------------------------------------------------------------
# .txt parser
# ---------------------------------------------------------------------------

def _parse_txt(data: bytes) -> str:
    """Parse plain text file with encoding detection."""
    try:
        import chardet
        detected = chardet.detect(data)
        encoding = detected.get("encoding") or "utf-8"
        confidence = detected.get("confidence", 0)
        logger.debug("Detected encoding: %s (confidence %.2f)", encoding, confidence)
    except ImportError:
        encoding = "utf-8"

    # Try detected encoding first, then fall back
    for enc in [encoding, "utf-8", "latin-1", "cp1252"]:
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue

    raise DocumentParseError(
        "Could not decode the text file. Please ensure it is UTF-8 encoded."
    )


# ---------------------------------------------------------------------------
# .docx parser
# ---------------------------------------------------------------------------

def _parse_docx(file_path: Path) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document
    except ImportError as e:
        raise DocumentParseError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from e

    try:
        doc = Document(str(file_path))
        paragraphs: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))

        if not paragraphs:
            raise DocumentParseError(
                "The Word document appears to be empty or contains no extractable text."
            )

        return "\n\n".join(paragraphs)

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"Failed to parse Word document: {e}") from e


# ---------------------------------------------------------------------------
# .pdf parser
# ---------------------------------------------------------------------------

def _parse_pdf_pymupdf(file_path: Path) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    import fitz  # PyMuPDF

    try:
        doc = fitz.open(str(file_path))
        pages: list[str] = []

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text")
            if text.strip():
                pages.append(text.strip())

        doc.close()

        if not pages:
            raise DocumentParseError(
                "Could not extract text from PDF. The file may be scanned or image-based."
            )

        return "\n\n".join(pages)

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"PyMuPDF failed to parse PDF: {e}") from e


def _parse_pdf_pdfplumber(file_path: Path) -> str:
    """Extract text from PDF using pdfplumber (fallback)."""
    import pdfplumber

    try:
        with pdfplumber.open(str(file_path)) as pdf:
            pages: list[str] = []
            for page in pdf.pages:
                text = page.extract_text()
                if text and text.strip():
                    pages.append(text.strip())

        if not pages:
            raise DocumentParseError(
                "Could not extract text from PDF. The file may be scanned or image-based."
            )

        return "\n\n".join(pages)

    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"pdfplumber failed to parse PDF: {e}") from e


def _parse_pdf(file_path: Path) -> str:
    """Parse PDF, trying PyMuPDF first, then pdfplumber."""
    try:
        return _parse_pdf_pymupdf(file_path)
    except ImportError:
        logger.info("PyMuPDF not available, falling back to pdfplumber")
    except DocumentParseError:
        logger.info("PyMuPDF extraction failed, trying pdfplumber")

    try:
        return _parse_pdf_pdfplumber(file_path)
    except ImportError as e:
        raise DocumentParseError(
            "Neither PyMuPDF nor pdfplumber is installed. "
            "Run: pip install PyMuPDF pdfplumber"
        ) from e


# ---------------------------------------------------------------------------
# .pptx parser
# ---------------------------------------------------------------------------

def _parse_pptx(file_path: Path) -> str:
    """Extract readable text from PowerPoint slides with slide provenance."""
    try:
        from pptx import Presentation
    except ImportError as e:
        raise DocumentParseError(
            "python-pptx is not installed. Run: pip install python-pptx"
        ) from e

    try:
        presentation = Presentation(str(file_path))
        slides: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if isinstance(text, str) and text.strip():
                    texts.append(text.strip())
            if texts:
                slides.append(f"[Slide {slide_number}]\n" + "\n".join(texts))
        if not slides:
            raise DocumentParseError("The PowerPoint contains no extractable text.")
        return "\n\n".join(slides)
    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"Failed to parse PowerPoint: {e}") from e


# ---------------------------------------------------------------------------
# .xlsx parser
# ---------------------------------------------------------------------------

def _parse_xlsx(file_path: Path) -> str:
    """Extract non-empty spreadsheet rows with worksheet provenance."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise DocumentParseError(
            "openpyxl is not installed. Run: pip install openpyxl"
        ) from e

    try:
        workbook = load_workbook(str(file_path), read_only=True, data_only=True)
        sheets: list[str] = []
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append(" | ".join(values))
            if rows:
                sheets.append(f"[Sheet: {worksheet.title}]\n" + "\n".join(rows))
        workbook.close()
        if not sheets:
            raise DocumentParseError("The Excel workbook contains no non-empty cells.")
        return "\n\n".join(sheets)
    except DocumentParseError:
        raise
    except Exception as e:
        raise DocumentParseError(f"Failed to parse Excel workbook: {e}") from e


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def parse_document(file_bytes: bytes, filename: str) -> str:
    """
    Parse an uploaded document and return normalized plain text.

    Args:
        file_bytes: Raw bytes of the uploaded file.
        filename:   Original filename (used to determine file type).

    Returns:
        Clean, normalized text extracted from the document.

    Raises:
        DocumentParseError: On any parsing failure.
    """
    from utils.file_utils import TempFileManager

    ext = Path(filename).suffix.lstrip(".").lower()
    logger.info("Parsing document: %s (ext=%s, size=%d bytes)", filename, ext, len(file_bytes))

    if not file_bytes:
        raise DocumentParseError("The uploaded file is empty.")

    # .txt — parse directly from bytes (no temp file needed)
    if ext == "txt":
        raw_text = _parse_txt(file_bytes)
        clean = clean_text(raw_text)
        logger.info("Extracted %d chars from .txt file", len(clean))
        return clean

    # Office documents — need a temp file
    if ext in ("docx", "pdf", "pptx", "xlsx"):
        with TempFileManager(file_bytes, suffix=f".{ext}") as tmp_path:
            if ext == "docx":
                raw_text = _parse_docx(tmp_path)
            elif ext == "pdf":
                raw_text = _parse_pdf(tmp_path)
            elif ext == "pptx":
                raw_text = _parse_pptx(tmp_path)
            else:
                raw_text = _parse_xlsx(tmp_path)

        clean = clean_text(raw_text)
        logger.info("Extracted %d chars from .%s file", len(clean), ext)
        return clean

    raise DocumentParseError(
        f"Unsupported file type '.{ext}'. Supported: txt, docx, pdf, pptx, xlsx"
    )


# ---------------------------------------------------------------------------
# Semantic-aware entry point (new pipeline)
# ---------------------------------------------------------------------------

def parse_document_with_structure(
    file_bytes: bytes,
    filename: str,
) -> tuple:
    """
    Parse an uploaded document and return BOTH the plain text AND a
    DocumentSection hierarchy tree for semantic chunking.

    Returns:
        (plain_text: str, root_section: DocumentSection, detection_method: str)

    The plain_text is identical to what parse_document() returns.
    The root_section contains the heading hierarchy (may have 0 subsections
    if the document has no detectable headings — caller handles this case).

    Raises:
        DocumentParseError: On any parsing failure.
    """
    from utils.file_utils import TempFileManager
    from core.document_structure import (
        extract_structure_from_docx,
        extract_structure_from_txt,
        extract_structure_from_pdf,
        DocumentSection,
    )

    ext = Path(filename).suffix.lstrip(".").lower()
    logger.info(
        "Parsing document with structure: %s (ext=%s, size=%d bytes)",
        filename, ext, len(file_bytes),
    )

    if not file_bytes:
        raise DocumentParseError("The uploaded file is empty.")

    if ext == "txt":
        raw_text = _parse_txt(file_bytes)
        plain_text = clean_text(raw_text)
        root, method = extract_structure_from_txt(plain_text)
        logger.info("TXT: %d chars, detection=%s", len(plain_text), method)
        return plain_text, root, method

    if ext in ("docx", "pdf", "pptx", "xlsx"):
        with TempFileManager(file_bytes, suffix=f".{ext}") as tmp_path:
            if ext == "docx":
                raw_text = _parse_docx(tmp_path)
                plain_text = clean_text(raw_text)
                root, method = extract_structure_from_docx(tmp_path)
            elif ext == "pdf":
                raw_text = _parse_pdf(tmp_path)
                plain_text = clean_text(raw_text)
                root, method = extract_structure_from_pdf(tmp_path)
            else:
                raw_text = _parse_pptx(tmp_path) if ext == "pptx" else _parse_xlsx(tmp_path)
                plain_text = clean_text(raw_text)
                root, method = extract_structure_from_txt(plain_text)

        logger.info("%s: %d chars, detection=%s", ext.upper(), len(plain_text), method)
        return plain_text, root, method

    raise DocumentParseError(
        f"Unsupported file type '.{ext}'. Supported: txt, docx, pdf, pptx, xlsx"
    )
