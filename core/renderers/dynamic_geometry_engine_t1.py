"""
core/renderers/dynamic_geometry_engine_t1.py — Dynamic slide layout geometry calculator & renderer
for Template-1 (AITransformationWeeklyUpdate4SEP2026.pptx).

Renders arbitrary LLM-generated layout compositions inside the Template-1 content canvas.
Matches the exact warm earth-tone palette, Georgia/Aptos typography, and design language
of the source template — zero deviation from brand style.
"""
from __future__ import annotations

from typing import Any, Dict, List, Optional
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from llm.dynamic_layout_schemas import (
    CardItem,
    ComparisonColumn,
    DynamicLayoutComposition,
    DynamicSlideDefinition,
    HeroBlock,
    KPIMetricItem,
    ProcessStepItem,
    TableDataGrid,
    TimelineItem,
)
from utils.logging_utils import get_logger

logger = get_logger(__name__)

# ---------------------------------------------------------------------------
# Template-1 Design System Tokens  (warm earth / sand palette)
# ---------------------------------------------------------------------------
COLOR_ESPRESSO       = RGBColor(0x35, 0x1C, 0x15)   # #351C15  header bar, labels
COLOR_AMBER_GOLD     = RGBColor(0xFF, 0xB5, 0x00)   # #FFB500  accent / KPI highlight
COLOR_STATUS_AMBER   = RGBColor(0xD4, 0x8A, 0x00)   # #D48A00  status pill / badge
COLOR_GREEN_POSITIVE = RGBColor(0x2E, 0x7D, 0x4F)   # #2E7D4F  positive metrics

COLOR_BG_SLIDE       = RGBColor(0xF7, 0xF4, 0xEF)   # #F7F4EF  slide background (warm sand)
COLOR_BG_CARD_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)   # #FFFFFF  primary card / row
COLOR_BG_CARD_BEIGE  = RGBColor(0xF2, 0xEC, 0xE5)   # #F2ECE5  alternating row / card
COLOR_BG_SECTION     = RGBColor(0xEE, 0xE7, 0xDE)   # #EEE7DE  section / leadership bar
COLOR_BG_HIGHLIGHT   = RGBColor(0xF7, 0xE7, 0xC1)   # #F7E7C1  amber tint callout

COLOR_BORDER         = RGBColor(0xE2, 0xD8, 0xCC)   # #E2D8CC  card/row border
COLOR_BORDER_INNER   = RGBColor(0xDD, 0xD6, 0xCE)   # #DDD6CE  inner grid lines

COLOR_TEXT_PRIMARY   = RGBColor(0x3A, 0x35, 0x31)   # #3A3531  body text
COLOR_TEXT_DARK      = RGBColor(0x35, 0x1C, 0x15)   # #351C15  labels / emphasis
COLOR_TEXT_MUTED     = RGBColor(0x79, 0x6B, 0x61)   # #796B61  captions / secondary
COLOR_TEXT_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)   # #FFFFFF  text on dark bg
COLOR_TEXT_CREAM     = RGBColor(0xFA, 0xF7, 0xF3)   # #FAF7F3  subtitle on dark bg

FONT_TITLE  = "Georgia"
FONT_BODY   = "Aptos"
FONT_FOOTER = "Calibri"


class DynamicGeometryEngineT1:
    """
    Renders dynamic visual layouts inside the 16:9 content canvas of Template-1.

    Safe canvas (below 1.11-in header bar, above footer):
        Left:   0.54"  (495300 EMU)
        Top:    1.25"  (1143000 EMU)
        Width:  12.25" (11201400 EMU)
        Height: 5.30"  (4846080 EMU)
    """

    DEFAULT_LEFT   = 495300
    DEFAULT_TOP    = 1143000
    DEFAULT_WIDTH  = 11201400
    DEFAULT_HEIGHT = 4846080

    def __init__(self, font_body: str = FONT_BODY, font_title: str = FONT_TITLE):
        self.font_body  = font_body
        self.font_title = font_title

    # -----------------------------------------------------------------------
    # Primary entry-point
    # -----------------------------------------------------------------------

    def render_slide_content(self, slide: Any, slide_def: DynamicSlideDefinition) -> None:
        """Renders the complete visual layout for a dynamic content slide."""
        curr_top    = self.DEFAULT_TOP
        curr_height = self.DEFAULT_HEIGHT

        if slide_def.executive_takeaway:
            banner_h = 400000
            self._render_executive_takeaway(
                slide,
                left=self.DEFAULT_LEFT, top=curr_top,
                width=self.DEFAULT_WIDTH, height=banner_h,
                text=slide_def.executive_takeaway,
            )
            curr_top    += banner_h + 160000
            curr_height -= (banner_h + 160000)

        bounds = {
            "left":   self.DEFAULT_LEFT,
            "top":    curr_top,
            "width":  self.DEFAULT_WIDTH,
            "height": curr_height,
        }

        comp    = slide_def.composition
        pattern = (comp.pattern or "").upper()

        logger.info("T1 Rendering slide %d — pattern: %s", slide_def.slide_number, pattern)

        if "METRIC_RIBBON" in pattern and comp.metrics:
            self._render_metric_ribbon_and_cards(slide, bounds, comp)
        elif "ACTION_TABLE" in pattern:
            self._render_action_table(slide, bounds, comp)
        elif "PROCESS" in pattern and comp.steps:
            self._render_process_pipeline(slide, bounds, comp.steps)
        elif "COMPARISON" in pattern and comp.comparison:
            self._render_comparison_split(slide, bounds, comp.comparison)
        elif "TABLE" in pattern and comp.table:
            self._render_data_matrix_table(slide, bounds, comp.table)
        elif "TIMELINE" in pattern and comp.timeline:
            self._render_timeline_roadmap(slide, bounds, comp.timeline)
        elif "HERO" in pattern and comp.hero:
            self._render_hero_and_sidebar(slide, bounds, comp)
        elif "KEY_HIGHLIGHTS" in pattern and comp.cards:
            self._render_key_highlights(slide, bounds, comp.cards)
        elif comp.cards:
            self._render_multi_column_cards(slide, bounds, comp.cards)
        else:
            self._render_fallback(slide, bounds, comp, slide_def)

    # -----------------------------------------------------------------------
    # Component Renderers
    # -----------------------------------------------------------------------

    def _render_executive_takeaway(
        self, slide: Any, left: int, top: int, width: int, height: int, text: str
    ) -> None:
        """Amber-tinted takeaway banner beneath the header bar."""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        if shape.adjustments:
            shape.adjustments[0] = 0.08
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_BG_HIGHLIGHT
        shape.line.color.rgb      = COLOR_STATUS_AMBER
        shape.line.width          = Pt(1.0)

        tf = shape.text_frame
        tf.word_wrap     = True
        tf.margin_left   = Inches(0.2)
        tf.margin_top    = Inches(0.07)
        tf.margin_right  = Inches(0.2)
        tf.margin_bottom = Inches(0.07)

        p       = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT

        run_tag              = p.add_run()
        run_tag.text         = "KEY TAKEAWAY:  "
        run_tag.font.name    = self.font_body
        run_tag.font.size    = Pt(11)
        run_tag.font.bold    = True
        run_tag.font.color.rgb = COLOR_ESPRESSO

        run_body              = p.add_run()
        run_body.text         = text
        run_body.font.name    = self.font_body
        run_body.font.size    = Pt(11)
        run_body.font.bold    = False
        run_body.font.color.rgb = COLOR_TEXT_PRIMARY

    def _render_metric_ribbon_and_cards(
        self, slide: Any, bounds: Dict[str, int], comp: DynamicLayoutComposition
    ) -> None:
        """Top KPI metric blocks (warm palette) + bottom content cards."""
        metrics  = comp.metrics or []
        cards    = comp.cards   or []
        ribbon_h = int(bounds["height"] * 0.32)
        gap      = 160000
        cards_top = bounds["top"] + ribbon_h + gap
        cards_h  = bounds["height"] - ribbon_h - gap

        n   = min(max(len(metrics), 1), 4)
        m_g = 140000
        m_w = (bounds["width"] - (n - 1) * m_g) // n

        for i, m in enumerate(metrics[:n]):
            m_left = bounds["left"] + i * (m_w + m_g)
            self._draw_kpi_card(
                slide, left=m_left, top=bounds["top"],
                width=m_w, height=ribbon_h,
                value=m.value, label=m.label, delta=m.delta,
            )

        if cards:
            self._render_multi_column_cards(
                slide,
                {"left": bounds["left"], "top": cards_top,
                 "width": bounds["width"], "height": cards_h},
                cards,
            )

    def _render_action_table(
        self, slide: Any, bounds: Dict[str, int], comp: DynamicLayoutComposition
    ) -> None:
        """
        Template-1 native tabular action-plan layout.

        Rows come from comp.cards:
          card.title      -> Priority label   (col 1)
          card.bullets[0] -> Action text      (col 2) — multi-line if bullets > 1
          card.tag        -> Status/Due badge (col 3)
          card.footer     -> Owner            (col 4)

        A leadership-ask footer bar is rendered at the bottom using
        comp.hero.headline / comp.hero.subtext (if provided).
        Falls back to comp.table if no cards supplied.
        """
        if not comp.cards and comp.table:
            self._render_data_matrix_table(slide, bounds, comp.table)
            return

        rows = comp.cards or []

        # Column proportions (matching template)
        col_p_left  = bounds["left"]
        col_p_w     = int(bounds["width"] * 0.185)
        col_a_left  = col_p_left + col_p_w + 40000
        col_a_w     = int(bounds["width"] * 0.430)
        col_s_left  = col_a_left + col_a_w + 40000
        col_s_w     = int(bounds["width"] * 0.135)
        col_o_left  = col_s_left + col_s_w + 40000
        col_o_w     = bounds["left"] + bounds["width"] - col_o_left

        header_h = 480000
        row_h    = 557784
        footer_h = 484632

        avail    = bounds["height"] - header_h - 40000
        n_rows   = min(len(rows), max(1, avail // (row_h + 10000)))

        # ── Header bar ──────────────────────────────────────────────────
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            bounds["left"], bounds["top"], bounds["width"], header_h,
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_ESPRESSO
        bar.line.fill.background()

        headers       = ["Priority", "Executive outcome / next action", "Due", "Owner"]
        h_lefts       = [col_p_left,  col_a_left,  col_s_left,  col_o_left]
        h_widths      = [col_p_w,     col_a_w,     col_s_w,     col_o_w]

        for hdr_text, h_left, h_width in zip(headers, h_lefts, h_widths):
            self._add_row_text_box(
                slide, h_left, bounds["top"], h_width, header_h,
                text=hdr_text, font_size=10.0, bold=True,
                color=COLOR_TEXT_WHITE, h_indent=0.10, v_indent=0.14,
            )

        # ── Data rows ───────────────────────────────────────────────────
        row_top = bounds["top"] + header_h

        for i, card in enumerate(rows[:n_rows]):
            bg = COLOR_BG_CARD_WHITE if i % 2 == 0 else COLOR_BG_CARD_BEIGE

            row_bg = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, bounds["left"], row_top, bounds["width"], row_h,
            )
            row_bg.fill.solid()
            row_bg.fill.fore_color.rgb = bg
            row_bg.line.color.rgb      = COLOR_BORDER
            row_bg.line.width          = Pt(0.5)

            # Priority
            self._add_row_text_box(
                slide, col_p_left, row_top, col_p_w, row_h,
                text=card.title, font_size=10.5, bold=True,
                color=COLOR_TEXT_DARK, h_indent=0.14,
            )

            # Action description
            action_text = "\n".join(card.bullets[:3]) if card.bullets else ""
            self._add_row_text_box(
                slide, col_a_left, row_top, col_a_w, row_h,
                text=action_text, font_size=10.0, bold=False,
                color=COLOR_TEXT_PRIMARY, h_indent=0.10,
            )

            # Status badge
            if card.tag:
                badge = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    col_s_left + 30000,
                    row_top + int((row_h - 292608) / 2),
                    col_s_w - 60000,
                    292608,
                )
                badge.fill.solid()
                badge.fill.fore_color.rgb = COLOR_STATUS_AMBER
                badge.line.color.rgb      = COLOR_STATUS_AMBER
                badge.line.width          = Pt(1.0)
                btf = badge.text_frame
                btf.margin_top = Inches(0.04)
                bp = btf.paragraphs[0]
                bp.alignment = PP_ALIGN.CENTER
                br = bp.add_run()
                br.text           = card.tag
                br.font.name      = self.font_body
                br.font.size      = Pt(8.5)
                br.font.bold      = True
                br.font.color.rgb = COLOR_TEXT_WHITE

            # Owner
            if card.footer:
                self._add_row_text_box(
                    slide, col_o_left, row_top, col_o_w, row_h,
                    text=card.footer, font_size=10.0, bold=False,
                    color=COLOR_TEXT_PRIMARY, h_indent=0.10,
                )

            row_top += row_h

        # ── Leadership footer bar ────────────────────────────────────────
        footer_top = min(
            row_top + 30000,
            bounds["top"] + bounds["height"] - footer_h,
        )
        footer_bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, bounds["left"], footer_top, bounds["width"], footer_h,
        )
        footer_bg.fill.solid()
        footer_bg.fill.fore_color.rgb = COLOR_BG_SECTION
        footer_bg.line.color.rgb      = COLOR_BORDER
        footer_bg.line.width          = Pt(0.6)

        self._add_row_text_box(
            slide,
            bounds["left"] + 100000, footer_top,
            int(bounds["width"] * 0.13), footer_h,
            text="Leadership ask", font_size=10.0, bold=True,
            color=COLOR_TEXT_DARK, h_indent=0.08,
        )

        ask_text = ""
        if comp.hero:
            ask_text = comp.hero.headline
            if comp.hero.subtext:
                ask_text += f" — {comp.hero.subtext}"
        elif rows:
            ask_text = "Resolve outstanding items and maintain delivery timelines."

        self._add_row_text_box(
            slide,
            bounds["left"] + int(bounds["width"] * 0.15), footer_top,
            int(bounds["width"] * 0.84), footer_h,
            text=ask_text, font_size=10.5, bold=False,
            color=COLOR_TEXT_PRIMARY, h_indent=0.08,
        )

    def _render_multi_column_cards(
        self, slide: Any, bounds: Dict[str, int], cards: List[CardItem]
    ) -> None:
        """2–4 vertical warm-palette content cards."""
        n   = min(max(len(cards), 1), 4)
        gap = 180000
        c_w = (bounds["width"] - (n - 1) * gap) // n

        max_bullets = max(len(c.bullets or []) for c in cards[:n]) if cards else 3
        needed_h = 1150000 + max_bullets * 420000 + (320000 if any(getattr(c, 'footer', None) for c in cards[:n]) else 0)
        card_h = min(bounds["height"], max(2600000, needed_h))
        card_top = bounds["top"] + max(0, (bounds["height"] - card_h) // 4)

        for i, card in enumerate(cards[:n]):
            c_left = bounds["left"] + i * (c_w + gap)
            self._draw_content_card(
                slide,
                left=c_left, top=card_top,
                width=c_w, height=card_h,
                title=card.title, tag=card.tag,
                bullets=card.bullets, footer=card.footer,
                alt_bg=(i % 2 == 1),
            )

    def _render_process_pipeline(
        self, slide: Any, bounds: Dict[str, int], steps: List[ProcessStepItem]
    ) -> None:
        """3–5 numbered step cards with amber step badges."""
        n   = min(max(len(steps), 1), 5)
        gap = 140000
        s_w = (bounds["width"] - (n - 1) * gap) // n

        # Adaptive height based on description length
        max_desc_len = max(len(s.description or "") for s in steps[:n]) if steps else 50
        needed_h = 1600000 + (max_desc_len // 30) * 280000
        card_h = min(bounds["height"], max(2700000, needed_h))
        card_top = bounds["top"] + max(0, (bounds["height"] - card_h) // 4)

        for i, step in enumerate(steps[:n]):
            s_left = bounds["left"] + i * (s_w + gap)

            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, s_left, card_top, s_w, card_h
            )
            if shape.adjustments:
                shape.adjustments[0] = 0.035
            shape.fill.solid()
            shape.fill.fore_color.rgb = COLOR_BG_CARD_WHITE if i % 2 == 0 else COLOR_BG_CARD_BEIGE
            shape.line.color.rgb      = COLOR_BORDER
            shape.line.width          = Pt(0.8)

            badge_h = 380000
            badge   = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                s_left + 110000, card_top + 110000,
                s_w - 220000, badge_h,
            )
            if badge.adjustments:
                badge.adjustments[0] = 0.12
            badge.fill.solid()
            badge.fill.fore_color.rgb = COLOR_ESPRESSO if i == 0 else COLOR_STATUS_AMBER
            badge.line.fill.background()

            btf = badge.text_frame
            btf.margin_top = Inches(0.04)
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            br = bp.add_run()
            br.text           = f"STEP {step.step_number}"
            br.font.name      = self.font_body
            br.font.size      = Pt(11)
            br.font.bold      = True
            br.font.color.rgb = COLOR_TEXT_WHITE

            tf = shape.text_frame
            tf.word_wrap    = True
            tf.margin_left  = Inches(0.20)
            tf.margin_right = Inches(0.20)
            tf.margin_top   = Inches(0.85)

            p_title = tf.paragraphs[0]
            r_title = p_title.add_run()
            r_title.text           = step.title
            r_title.font.name      = self.font_body
            r_title.font.size      = Pt(13.5)
            r_title.font.bold      = True
            r_title.font.color.rgb = COLOR_TEXT_DARK

            if step.tag:
                p_tag = tf.add_paragraph()
                p_tag.space_before = Pt(4)
                r_tag = p_tag.add_run()
                r_tag.text           = f"[{step.tag}]"
                r_tag.font.name      = self.font_body
                r_tag.font.size      = Pt(9.5)
                r_tag.font.bold      = True
                r_tag.font.color.rgb = COLOR_STATUS_AMBER

            p_desc = tf.add_paragraph()
            p_desc.space_before = Pt(8)
            r_desc = p_desc.add_run()
            r_desc.text           = step.description
            r_desc.font.name      = self.font_body
            r_desc.font.size      = Pt(11)
            r_desc.font.color.rgb = COLOR_TEXT_PRIMARY

    def _render_comparison_split(
        self, slide: Any, bounds: Dict[str, int], cols: List[ComparisonColumn]
    ) -> None:
        """2–3 side-by-side contrast panels."""
        n   = min(max(len(cols), 1), 3)
        gap = 220000
        c_w = (bounds["width"] - (n - 1) * gap) // n

        # Adaptive height based on content
        max_bullets = max(len(col.bullets or []) for col in cols[:n]) if cols else 3
        needed_h = 1150000 + max_bullets * 450000
        card_h = min(bounds["height"], max(2700000, needed_h))
        card_top = bounds["top"] + max(0, (bounds["height"] - card_h) // 4)

        for i, col in enumerate(cols[:n]):
            c_left    = bounds["left"] + i * (c_w + gap)
            is_target = i == len(cols) - 1 and len(cols) > 1
            self._draw_content_card(
                slide,
                left=c_left, top=card_top,
                width=c_w, height=card_h,
                title=col.heading,
                tag=col.tag or ("Target State" if is_target else "Current State"),
                bullets=col.bullets,
                highlight_border=is_target,
                alt_bg=(i % 2 == 1),
            )

    def _render_data_matrix_table(
        self, slide: Any, bounds: Dict[str, int], table_data: TableDataGrid
    ) -> None:
        """Warm-palette structured data table."""
        headers = table_data.headers or []
        rows    = table_data.rows    or []
        if not headers or not rows:
            return

        num_rows = min(len(rows) + 1, 10)
        num_cols = min(len(headers), 7)

        tbl = slide.shapes.add_table(
            num_rows, num_cols,
            bounds["left"], bounds["top"],
            bounds["width"], bounds["height"],
        ).table

        for c, hdr in enumerate(headers[:num_cols]):
            cell = tbl.cell(0, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_ESPRESSO
            cell.text = hdr
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for r in p.runs:
                    r.font.name      = self.font_body
                    r.font.size      = Pt(10.5)
                    r.font.bold      = True
                    r.font.color.rgb = COLOR_TEXT_WHITE

        for ri, row in enumerate(rows[:num_rows - 1]):
            bg = COLOR_BG_CARD_WHITE if ri % 2 == 0 else COLOR_BG_CARD_BEIGE
            for c in range(num_cols):
                cell = tbl.cell(ri + 1, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                val = row[c] if c < len(row) else ""
                cell.text = str(val)
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for r in p.runs:
                        r.font.name      = self.font_body
                        r.font.size      = Pt(10)
                        r.font.color.rgb = COLOR_TEXT_PRIMARY

    def _render_timeline_roadmap(
        self, slide: Any, bounds: Dict[str, int], timeline: List[TimelineItem]
    ) -> None:
        """3–5 chronological milestone cards with amber date pills."""
        n   = min(max(len(timeline), 1), 5)
        gap = 130000
        t_w = (bounds["width"] - (n - 1) * gap) // n

        for i, item in enumerate(timeline[:n]):
            t_left = bounds["left"] + i * (t_w + gap)

            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, t_left, bounds["top"], t_w, bounds["height"]
            )
            if shape.adjustments:
                shape.adjustments[0] = 0.035
            shape.fill.solid()
            shape.fill.fore_color.rgb = COLOR_BG_CARD_WHITE if i % 2 == 0 else COLOR_BG_CARD_BEIGE
            shape.line.color.rgb      = COLOR_BORDER
            shape.line.width          = Pt(0.8)

            pill_h = 340000
            pill   = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                t_left + 110000, bounds["top"] + 110000,
                t_w - 220000, pill_h,
            )
            if pill.adjustments:
                pill.adjustments[0] = 0.2
            pill.fill.solid()
            pill.fill.fore_color.rgb = COLOR_STATUS_AMBER
            pill.line.fill.background()

            ptf = pill.text_frame
            ptf.margin_top = Inches(0.04)
            pp = ptf.paragraphs[0]
            pp.alignment = PP_ALIGN.CENTER
            pr = pp.add_run()
            pr.text           = item.date
            pr.font.name      = self.font_body
            pr.font.size      = Pt(10.5)
            pr.font.bold      = True
            pr.font.color.rgb = COLOR_TEXT_WHITE

            tf = shape.text_frame
            tf.word_wrap    = True
            tf.margin_left  = Inches(0.18)
            tf.margin_right = Inches(0.18)
            tf.margin_top   = Inches(0.72)

            p_title = tf.paragraphs[0]
            r_title = p_title.add_run()
            r_title.text           = item.event
            r_title.font.name      = self.font_body
            r_title.font.size      = Pt(12)
            r_title.font.bold      = True
            r_title.font.color.rgb = COLOR_TEXT_DARK

            p_desc = tf.add_paragraph()
            p_desc.space_before = Pt(8)
            r_desc = p_desc.add_run()
            r_desc.text           = item.description
            r_desc.font.name      = self.font_body
            r_desc.font.size      = Pt(10)
            r_desc.font.color.rgb = COLOR_TEXT_MUTED

    def _render_hero_and_sidebar(
        self, slide: Any, bounds: Dict[str, int], comp: DynamicLayoutComposition
    ) -> None:
        """1/3 espresso-dark sidebar hero + 2/3 warm content cards."""
        hero  = comp.hero
        cards = comp.cards or []

        sidebar_w = int(bounds["width"] * 0.30)
        gap       = 160000
        main_left = bounds["left"] + sidebar_w + gap
        main_w    = bounds["width"] - sidebar_w - gap

        if hero:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                bounds["left"], bounds["top"], sidebar_w, bounds["height"],
            )
            if shape.adjustments:
                shape.adjustments[0] = 0.035
            shape.fill.solid()
            shape.fill.fore_color.rgb = COLOR_ESPRESSO
            shape.line.fill.background()

            tf = shape.text_frame
            tf.word_wrap    = True
            tf.margin_left  = Inches(0.25)
            tf.margin_right = Inches(0.25)
            tf.margin_top   = Inches(0.35)

            p_tag = tf.paragraphs[0]
            r_tag = p_tag.add_run()
            r_tag.text           = "EXECUTIVE BRIEF"
            r_tag.font.name      = self.font_body
            r_tag.font.size      = Pt(10)
            r_tag.font.bold      = True
            r_tag.font.color.rgb = COLOR_AMBER_GOLD

            p_head = tf.add_paragraph()
            p_head.space_before = Pt(10)
            r_head = p_head.add_run()
            r_head.text           = hero.headline
            r_head.font.name      = self.font_title
            r_head.font.size      = Pt(16)
            r_head.font.bold      = True
            r_head.font.color.rgb = COLOR_TEXT_WHITE

            p_sub = tf.add_paragraph()
            p_sub.space_before = Pt(10)
            r_sub = p_sub.add_run()
            r_sub.text           = hero.subtext
            r_sub.font.name      = self.font_body
            r_sub.font.size      = Pt(11)
            r_sub.font.color.rgb = COLOR_TEXT_CREAM

        if cards:
            self._render_multi_column_cards(
                slide,
                {"left": main_left, "top": bounds["top"],
                 "width": main_w, "height": bounds["height"]},
                cards,
            )

    def _render_key_highlights(
        self, slide: Any, bounds: Dict[str, int], cards: List[CardItem]
    ) -> None:
        self._render_multi_column_cards(slide, bounds, cards)

    def _render_fallback(
        self,
        slide: Any,
        bounds: Dict[str, int],
        comp: DynamicLayoutComposition,
        slide_def: Optional[DynamicSlideDefinition] = None,
    ) -> None:
        """Context-aware fallback renderer using actual slide content."""
        title_text = slide_def.title if slide_def else "Strategic Overview"
        takeaway_text = (slide_def.executive_takeaway if slide_def else "") or "Core strategic priorities and execution highlights."

        self._draw_content_card(
            slide,
            left=bounds["left"],
            top=bounds["top"],
            width=bounds["width"],
            height=bounds["height"],
            title=title_text,
            tag="KEY FOCUS",
            bullets=[takeaway_text],
            alt_bg=False,
        )

    # -----------------------------------------------------------------------
    # Primitive Drawing Helpers
    # -----------------------------------------------------------------------

    def _draw_kpi_card(
        self,
        slide: Any,
        left: int, top: int, width: int, height: int,
        value: str, label: str, delta: Optional[str] = None,
    ) -> None:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        if shape.adjustments:
            shape.adjustments[0] = 0.035
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_BG_CARD_WHITE
        shape.line.color.rgb      = COLOR_BORDER
        shape.line.width          = Pt(0.8)

        tf = shape.text_frame
        tf.word_wrap    = True
        tf.margin_left  = Inches(0.22)
        tf.margin_right = Inches(0.22)
        tf.margin_top   = Inches(0.18)

        p_val = tf.paragraphs[0]
        p_val.alignment = PP_ALIGN.LEFT
        r_val = p_val.add_run()
        r_val.text           = value
        r_val.font.name      = self.font_body
        r_val.font.size      = Pt(28)
        r_val.font.bold      = True
        r_val.font.color.rgb = COLOR_AMBER_GOLD

        p_lbl = tf.add_paragraph()
        p_lbl.space_before = Pt(3)
        r_lbl = p_lbl.add_run()
        r_lbl.text           = label
        r_lbl.font.name      = self.font_body
        r_lbl.font.size      = Pt(11.5)
        r_lbl.font.bold      = True
        r_lbl.font.color.rgb = COLOR_TEXT_DARK

        if delta:
            p_del = tf.add_paragraph()
            p_del.space_before = Pt(3)
            r_del = p_del.add_run()
            r_del.text           = delta
            r_del.font.name      = self.font_body
            r_del.font.size      = Pt(10)
            r_del.font.color.rgb = COLOR_TEXT_MUTED

    def _draw_content_card(
        self,
        slide: Any,
        left: int, top: int, width: int, height: int,
        title: str,
        tag: Optional[str] = None,
        bullets: Optional[List[str]] = None,
        footer: Optional[str] = None,
        highlight_border: bool = False,
        alt_bg: bool = False,
    ) -> None:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        if shape.adjustments:
            shape.adjustments[0] = 0.035
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_BG_CARD_BEIGE if alt_bg else COLOR_BG_CARD_WHITE
        shape.line.color.rgb      = COLOR_STATUS_AMBER if highlight_border else COLOR_BORDER
        shape.line.width          = Pt(1.5 if highlight_border else 0.8)

        tf = shape.text_frame
        tf.word_wrap    = True
        tf.margin_left  = Inches(0.28)
        tf.margin_right = Inches(0.28)
        tf.margin_top   = Inches(0.24)
        tf.margin_bottom = Inches(0.24)

        if tag:
            p_tag = tf.paragraphs[0]
            r_tag = p_tag.add_run()
            r_tag.text           = tag.upper()
            r_tag.font.name      = self.font_body
            r_tag.font.size      = Pt(9.5)
            r_tag.font.bold      = True
            r_tag.font.color.rgb = COLOR_STATUS_AMBER
            p_title = tf.add_paragraph()
            p_title.space_before = Pt(4)
        else:
            p_title = tf.paragraphs[0]

        r_title = p_title.add_run()
        r_title.text           = title
        r_title.font.name      = self.font_body
        r_title.font.size      = Pt(14)
        r_title.font.bold      = True
        r_title.font.color.rgb = COLOR_TEXT_DARK

        b_list = bullets or []
        n_bullets = len(b_list)
        b_font_size = Pt(12) if n_bullets <= 3 else Pt(11)
        b_space = Pt(10) if n_bullets <= 3 else Pt(7)

        for bullet_text in b_list:
            p_b = tf.add_paragraph()
            p_b.space_before = b_space
            r_b = p_b.add_run()
            r_b.text           = f"\u2022  {bullet_text}"
            r_b.font.name      = self.font_body
            r_b.font.size      = b_font_size
            r_b.font.color.rgb = COLOR_TEXT_PRIMARY

        if footer:
            p_foot = tf.add_paragraph()
            p_foot.space_before = Pt(10)
            r_foot = p_foot.add_run()
            r_foot.text           = footer
            r_foot.font.name      = self.font_body
            r_foot.font.size      = Pt(10)
            r_foot.font.italic    = True
            r_foot.font.color.rgb = COLOR_TEXT_MUTED

    def _add_row_text_box(
        self,
        slide: Any,
        left: int, top: int, width: int, height: int,
        text: str, font_size: float, bold: bool,
        color: RGBColor, h_indent: float = 0.12,
        v_indent: float = 0.10,
    ) -> None:
        """Transparent text shape for action-table row cells."""
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.background()
        shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap   = True
        tf.margin_left = Inches(h_indent)
        tf.margin_top  = Inches(v_indent)
        first = True
        for line in text.split("\n"):
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            r = p.add_run()
            r.text           = line
            r.font.name      = self.font_body
            r.font.size      = Pt(font_size)
            r.font.bold      = bold
            r.font.color.rgb = color
