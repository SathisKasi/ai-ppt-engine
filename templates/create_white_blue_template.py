"""
templates/create_white_blue_template.py
Generates the White+Blue professional presentation template.
Run: python templates/create_white_blue_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLUE_PRIMARY  = RGBColor(0x00, 0x57, 0xB8)
BLUE_DEEP     = RGBColor(0x00, 0x3D, 0x82)
BLUE_LIGHT    = RGBColor(0x00, 0xA0, 0xE4)
BLUE_CARD     = RGBColor(0xF0, 0xF7, 0xFF)
BLUE_DIVIDER  = RGBColor(0xD0, 0xE8, 0xF8)
BLUE_ACCENT2  = RGBColor(0xE8, 0xF4, 0xFD)
TEXT_DARK     = RGBColor(0x1A, 0x1A, 0x2A)
TEXT_MED      = RGBColor(0x4A, 0x55, 0x68)
TEXT_MUTED    = RGBColor(0x94, 0xA3, 0xB8)
TEXT_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

def _bg(slide, color=BG_WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def _rect(slide, l, t, w, h, fill=BLUE_PRIMARY, lc=None):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if lc: s.line.color.rgb = lc; s.line.width = Pt(0.5)
    else:   s.line.fill.background()
    return s

def _txt(slide, text, l, t, w, h, sz=18, bold=False, color=TEXT_DARK,
         align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = str(text)
    run.font.size = Pt(sz); run.font.bold = bold; run.font.italic = italic
    run.font.color.rgb = color; run.font.name = FONT
    return tb

def _bullets(slide, items, l, t, w, h, sz=16, color=TEXT_MED):
    import re
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        cleaned = re.sub(r"^[\s\u2022\-\*]+", "", str(item)).strip()
        if not cleaned: continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run(); run.text = f"\u25cf  {cleaned}"
        run.font.size = Pt(sz); run.font.color.rgb = color; run.font.name = FONT
    return tb

def _hdr(slide, title, sz=26):
    _rect(slide, 0, 0, 13.333, 1.3, BLUE_PRIMARY)
    _rect(slide, 0, 1.3, 13.333, 0.04, BLUE_LIGHT)
    _txt(slide, title, 0.35, 0.1, 12.5, 1.0, sz=sz, bold=True, color=TEXT_WHITE)

def _ftr(slide):
    _rect(slide, 0, 7.44, 13.333, 0.04, BLUE_DIVIDER)

def _card(slide, l, t, w, h):
    return _rect(slide, l, t, w, h, BLUE_CARD, BLUE_DIVIDER)

# Layout makers
def mk_title(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s)
    _rect(s, 0, 0, 13.333, 0.05, BLUE_PRIMARY)
    _rect(s, 0, 0.05, 5.5, 7.45, BLUE_PRIMARY)
    _rect(s, 5.5, 0.05, 0.04, 7.45, BLUE_LIGHT)
    _txt(s, "PRESENTATION TITLE", 0.4, 2.5, 4.8, 1.8, sz=36, bold=True,
         color=TEXT_WHITE, align=PP_ALIGN.LEFT)
    _txt(s, "Subtitle or tagline goes here", 0.4, 4.4, 4.8, 0.6, sz=16,
         color=RGBColor(0xC0,0xDE,0xFF))
    _txt(s, "Presented by", 6.2, 2.9, 6.5, 0.4, sz=12, color=TEXT_MUTED, italic=True)
    _txt(s, "Author Name", 6.2, 3.3, 6.5, 0.5, sz=18, bold=True, color=TEXT_DARK)
    _txt(s, "Organisation  \u2022  Date", 6.2, 3.85, 6.5, 0.4, sz=13, color=TEXT_MED)
    _rect(s, 0, 7.3, 13.333, 0.2, BLUE_DEEP)

def mk_title_content(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Slide Title")
    _card(s, 0.4, 1.5, 12.55, 5.65)
    _bullets(s, ["Bullet point one", "Bullet point two", "Bullet point three"],
             0.65, 1.7, 12.1, 5.2); _ftr(s)

def mk_two_column(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Two Column Layout")
    _card(s, 0.4, 1.5, 5.95, 5.65)
    _rect(s, 0.4, 1.5, 5.95, 0.45, BLUE_PRIMARY)
    _txt(s, "Left Column", 0.55, 1.55, 5.65, 0.38, sz=14, bold=True, color=TEXT_WHITE)
    _bullets(s, ["Point 1", "Point 2"], 0.6, 2.1, 5.5, 4.8)
    _card(s, 6.95, 1.5, 5.95, 5.65)
    _rect(s, 6.95, 1.5, 5.95, 0.45, BLUE_DEEP)
    _txt(s, "Right Column", 7.1, 1.55, 5.65, 0.38, sz=14, bold=True, color=TEXT_WHITE)
    _bullets(s, ["Point 1", "Point 2"], 7.1, 2.1, 5.5, 4.8)
    _rect(s, 6.6, 1.5, 0.04, 5.65, BLUE_DIVIDER); _ftr(s)

def mk_section_header(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s, BLUE_ACCENT2)
    _rect(s, 0, 0, 13.333, 0.06, BLUE_PRIMARY)
    _rect(s, 0, 7.44, 13.333, 0.06, BLUE_PRIMARY)
    _rect(s, 1.5, 3.4, 10.333, 0.05, BLUE_PRIMARY)
    _txt(s, "SECTION", 1.5, 2.2, 10.333, 0.5, sz=13, bold=True,
         color=BLUE_PRIMARY, align=PP_ALIGN.CENTER)
    _txt(s, "Section Title Goes Here", 1.0, 2.7, 11.333, 1.4, sz=40, bold=True,
         color=TEXT_DARK, align=PP_ALIGN.CENTER)
    _txt(s, "Optional section description", 2.0, 3.6, 9.333, 0.6, sz=16,
         color=TEXT_MED, align=PP_ALIGN.CENTER)

def mk_image_text(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Image & Content")
    _card(s, 0.4, 1.5, 5.8, 5.65); _rect(s, 0.4, 1.5, 5.8, 0.04, BLUE_PRIMARY)
    _txt(s, "[ Image / Diagram ]", 0.5, 3.6, 5.6, 0.9, sz=15,
         color=TEXT_MUTED, align=PP_ALIGN.CENTER, italic=True)
    _card(s, 6.6, 1.5, 6.35, 5.65)
    _txt(s, "Content heading", 6.75, 1.65, 6.05, 0.5, sz=16, bold=True, color=BLUE_PRIMARY)
    _bullets(s, ["Key point one", "Key point two"], 6.75, 2.3, 6.05, 4.5); _ftr(s)

def mk_comparison(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Comparison")
    _rect(s, 0.4, 1.5, 6.0, 0.55, BLUE_PRIMARY)
    _txt(s, "Option A", 0.6, 1.57, 5.6, 0.42, sz=17, bold=True, color=TEXT_WHITE)
    _card(s, 0.4, 2.05, 6.0, 5.1); _bullets(s, ["Point 1","Point 2"], 0.6, 2.2, 5.6, 4.7)
    _txt(s, "VS", 6.47, 3.7, 0.4, 0.6, sz=12, bold=True, color=BLUE_DEEP, align=PP_ALIGN.CENTER)
    _rect(s, 6.95, 1.5, 6.0, 0.55, BLUE_DEEP)
    _txt(s, "Option B", 7.1, 1.57, 5.6, 0.42, sz=17, bold=True, color=TEXT_WHITE)
    _card(s, 6.95, 2.05, 6.0, 5.1); _bullets(s, ["Point 1","Point 2"], 7.1, 2.2, 5.6, 4.7)
    _ftr(s)

def mk_timeline(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Timeline / Roadmap")
    _rect(s, 0.5, 4.0, 12.333, 0.06, BLUE_PRIMARY)
    for i, (x, lbl, ev) in enumerate([(1.0,"Q1","Event 1"),(4.1,"Q2","Event 2"),
                                        (7.2,"Q3","Event 3"),(10.3,"Q4","Event 4")]):
        _rect(s, x+0.7, 3.8, 0.4, 0.4, BLUE_PRIMARY if i%2==0 else BLUE_DEEP)
        _txt(s, lbl, x+0.55, 4.3, 0.7, 0.4, sz=12, bold=True,
             color=BLUE_PRIMARY, align=PP_ALIGN.CENTER)
        ct = 2.0 if i%2==0 else 4.9
        _card(s, x+0.1, ct, 2.2, 1.5)
        _txt(s, ev, x+0.2, ct+0.1, 2.0, 1.2, sz=13, color=TEXT_DARK, align=PP_ALIGN.CENTER)
    _ftr(s)

def mk_process(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Process / Workflow")
    cols = [BLUE_PRIMARY, BLUE_DEEP, BLUE_LIGHT, RGBColor(0x5B,0xA3,0xD8)]
    for i in range(4):
        x = 0.5 + i*3.15; c = cols[i]
        _card(s, x, 2.0, 2.7, 4.7); _rect(s, x+0.85, 1.7, 1.0, 1.0, c)
        _txt(s, str(i+1), x+0.85, 1.7, 1.0, 1.0, sz=22, bold=True,
             color=TEXT_WHITE, align=PP_ALIGN.CENTER)
        _txt(s, f"Step {i+1}", x+0.1, 2.9, 2.5, 0.55, sz=14, bold=True,
             color=c, align=PP_ALIGN.CENTER)
        _txt(s, "Description text", x+0.15, 3.55, 2.4, 2.9, sz=12,
             color=TEXT_MED, align=PP_ALIGN.CENTER)
        if i < 3:
            _txt(s, "\u2192", x+2.68, 3.9, 0.5, 0.6, sz=22, bold=True,
                 color=BLUE_LIGHT, align=PP_ALIGN.CENTER)
    _ftr(s)

def mk_architecture(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "System Architecture")
    lcs = [BLUE_PRIMARY, BLUE_DEEP, BLUE_LIGHT, RGBColor(0x5B,0xA3,0xD8)]
    rows = [("Presentation Layer","UI, Web, Mobile"),("Application Layer","Business Logic, APIs"),
            ("Data Layer","Databases, Storage"),("Infrastructure Layer","Cloud, Networking")]
    for i, (nm, desc) in enumerate(rows):
        y = 1.55 + i*1.28; c = lcs[i]
        _rect(s, 0.4, y, 12.5, 1.1, c)
        _txt(s, nm, 0.6, y+0.06, 5.5, 0.5, sz=16, bold=True, color=TEXT_WHITE)
        _txt(s, desc, 6.2, y+0.06, 6.5, 0.5, sz=14, color=RGBColor(0xC8,0xE8,0xFF))
        if i<3: _txt(s, "\u2195", 6.4, y+1.1, 0.5, 0.22, sz=11,
                     color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    _ftr(s)

def mk_statistics(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s); _hdr(s, "Key Metrics")
    scs = [BLUE_PRIMARY, BLUE_DEEP, BLUE_LIGHT, RGBColor(0x5B,0xA3,0xD8)]
    for i, (v, lbl) in enumerate([("85%","Metric One"),("2.4x","Metric Two"),
                                   ("$4M","Metric Three"),("98%","Metric Four")]):
        x = 0.45 + i*3.2; c = scs[i]
        _card(s, x, 1.8, 2.85, 5.1); _rect(s, x, 1.8, 2.85, 0.08, c)
        _txt(s, v, x+0.1, 2.6, 2.65, 1.5, sz=52, bold=True, color=c, align=PP_ALIGN.CENTER)
        _txt(s, lbl, x+0.1, 4.3, 2.65, 2.0, sz=15, color=TEXT_MED, align=PP_ALIGN.CENTER)
    _ftr(s)

def mk_conclusion(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); _bg(s)
    _rect(s, 0, 0, 13.333, 2.0, BLUE_PRIMARY)
    _txt(s, "Conclusion & Next Steps", 0.4, 0.3, 12.5, 1.4, sz=30, bold=True, color=TEXT_WHITE)
    _card(s, 0.4, 2.2, 7.6, 5.0); _rect(s, 0.4, 2.2, 7.6, 0.35, BLUE_LIGHT)
    _txt(s, "KEY TAKEAWAYS", 0.6, 2.24, 5.0, 0.28, sz=11, bold=True, color=TEXT_WHITE)
    _bullets(s, ["Summary point 1","Summary point 2","Summary point 3"], 0.6, 2.65, 7.2, 4.3)
    _rect(s, 8.35, 2.2, 4.6, 5.0, BLUE_DEEP)
    _txt(s, "NEXT STEPS", 8.5, 2.3, 4.3, 0.4, sz=11, bold=True, color=RGBColor(0xC0,0xDE,0xFF))
    _txt(s, "1. First action\n2. Second action\n3. Third action",
         8.5, 2.8, 4.2, 3.8, sz=15, color=TEXT_WHITE)
    _rect(s, 3.0, 7.1, 7.333, 0.3, BLUE_LIGHT)
    _txt(s, "Thank You \u2014 Questions Welcome", 3.0, 7.1, 7.333, 0.3,
         sz=13, bold=True, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

def create_white_blue_template(output_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    makers = [mk_title, mk_title_content, mk_two_column, mk_section_header,
              mk_image_text, mk_comparison, mk_timeline, mk_process,
              mk_architecture, mk_statistics, mk_conclusion]
    names  = ["TITLE","TITLE_AND_CONTENT","TWO_COLUMN","SECTION_HEADER","IMAGE_TEXT",
              "COMPARISON","TIMELINE","PROCESS","ARCHITECTURE","STATISTICS","CONCLUSION"]
    for i, (maker, nm) in enumerate(zip(makers, names)):
        maker(prs)
        print(f"  [{i}] {nm}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"Saved: {output_path} ({output_path.stat().st_size//1024} KB)")

if __name__ == "__main__":
    out = PROJECT_ROOT / "templates" / "white_blue_template.pptx"
    create_white_blue_template(out)
